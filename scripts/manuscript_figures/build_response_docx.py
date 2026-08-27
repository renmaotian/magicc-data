#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Build ``response_to_reviewers.docx`` by rendering ``response_to_reviewers.md``.

WHAT CHANGED, AND WHY
---------------------
The previous version of this script did not read ``response_to_reviewers.md`` --
it *wrote* it.  Roughly 3,000 lines of response prose were hard-coded here as
Python data structures, ``build_markdown()`` serialised them to Markdown and
``main()`` did ``OUT_MD.write_text(build_markdown())`` before rendering the same
structures to Word.  The Markdown on disk has since been edited directly and
extensively (cross-reference renumbering across five figures and one table, a new
opening note, three factual corrections, 865 decorative bold spans stripped, 14
corrected Methods subsection titles), so re-running the old script would have
silently destroyed all of it.  That was recorded as trap T12 in
`the internal project log` §11.9.

This version inverts the dependency.  ``response_to_reviewers.md`` is now the
single source of truth and this script only *renders* it.  The hard-coded corpus
is gone, and so is the verbatim-slicing machinery that used to cut reviewer
comments out of ``1st_submission/reviewers_comments.txt`` at build time -- the
comments now live in the Markdown, where the verification pass below checks them
line by line instead.

Document furniture is cloned from ``1st_submission/manuscript_v4.docx`` rather
than accepting python-docx's US Letter / Calibri defaults, exactly as
``build_supplementary_docx.py`` now does: ``styles.xml`` (Title / Heading 1-3 /
Normal, Arial 11 pt), the A4 ``sectPr`` with 1 in margins, the header and footer
parts, ``theme``, ``fontTable``, ``numbering`` and ``settings``, and the
``w:tblPr`` / row templates of the v4 table, which are cloned for every one of
the ten Markdown tables.  The v4 body and its four images are discarded, so
nothing of the manuscript is inherited by accident and the file stays small.

WHAT THE RENDERING PRESERVES
----------------------------
  * reviewer and editor comments set in their own distinct block -- indented,
    shaded, ruled on the left, italic, 10.5 pt -- reproduced verbatim;
  * each comment followed by its ``Response.`` and ``Changes made.`` blocks;
  * Markdown ``#`` -> Title, ``##`` -> Heading 1, ``###`` -> Heading 2, which is
    the mapping the previous build used;
  * a thematic break (``---``) starts a new page once the section it closes has
    subsections of its own, and is set as a rule otherwise; that reproduces the
    previous layout, in which Reviewer 1, Reviewer 2, the editorial requirements
    and the closing section each begin on a fresh page while the title block and
    the revision note stay with the cover note;
  * tables laid out so they cannot overflow the text column.

VERIFICATION (all assertions; the script exits non-zero on any failure)
----------------------------------------------------------------------
reopen the file; diff the whole text stream of the ``.docx`` against the Markdown
block by block; count the 56 entries, the 62 ``###`` headings and the 72 verbatim
quote lines and compare the quote lines character by character; audit every bold
run against the set of bold spans the Markdown actually asks for; confirm nothing
is painted past the right margin; render to PDF and report the page count and
page size.

Run:  /path/to/anaconda3/bin/python build_response_docx.py
"""

from __future__ import annotations

import copy
import os
import re
import shutil
import subprocess
import sys
import tempfile

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

# --------------------------------------------------------------------------
# paths
# --------------------------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
RESUB = os.path.dirname(HERE)
NC = os.path.dirname(RESUB)
V4_DOCX = os.path.join(NC, "1st_submission", "manuscript_v4.docx")
MD = os.path.join(RESUB, "response_to_reviewers.md")
OUT = os.path.join(RESUB, "response_to_reviewers.docx")

# what the document must contain; asserted, not assumed
N_ENTRIES = 56          # comment blocks: quote + Response. + Changes made.
N_H3 = 62               # '###' headings: 56 entries + 6 cover-note sections
N_QUOTE_LINES = 72      # '>' lines carrying the reviewer / editor comments ...
N_QUOTE_PARAS = 64      # ... of which 8 are the blank '>' paragraph separators
                        # inside the four multi-paragraph editorial quotations

# --------------------------------------------------------------------------
# page geometry (A4, 1 in margins) -- read back from the file, not assumed
# --------------------------------------------------------------------------
USABLE_TWIPS = None          # filled in from sectPr
TABLE_TOTAL_TWIPS = 9000     # fits the 9,026-twip A4 text column

# quote-block appearance, carried over from the previous build
QUOTE_FILL = "F2F2F2"
QUOTE_BAR = "8C8C8C"
QUOTE_PT = 10.5
QUOTE_COLOR = RGBColor(0x33, 0x33, 0x33)
QUOTE_LEFT_TWIPS = 504       # 0.35 in
QUOTE_RIGHT_TWIPS = 216      # 0.15 in

# labels that open a block and should not be left stranded at a page foot
LABELS = ("Response.", "Changes made.")


# ==========================================================================
# markdown parsing
# ==========================================================================

def split_row(line: str) -> list:
    """Split a markdown table row on unescaped pipes."""
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    out, buf, i = [], [], 0
    while i < len(line):
        ch = line[i]
        if ch == "\\" and i + 1 < len(line) and line[i + 1] == "|":
            buf.append("|")
            i += 2
            continue
        if ch == "|":
            out.append("".join(buf).strip())
            buf = []
            i += 1
            continue
        buf.append(ch)
        i += 1
    out.append("".join(buf).strip())
    return out


SEP_RE = re.compile(r"^\|[\s:\-|]+\|?\s*$")
NUM_ITEM_RE = re.compile(r"^(\d+)\.\s+(.*)$")
BULLET_RE = re.compile(r"^\s*-\s+")


def parse_markdown(path: str) -> list:
    """Return a flat list of blocks: ``(kind, payload)``.

    Kinds: ``h`` (level, text), ``p`` (text), ``bullet`` (text),
    ``num`` (label, text), ``quote`` (list of verbatim paragraph strings),
    ``table`` (list of rows), ``hr`` (None).

    A blockquote line is *not* inline-parsed.  Reviewer and editor comments are
    reproduced verbatim, and several of them legitimately begin with the
    editor's own ``*`` bullet character, which an inline parser would eat.  The
    outer ``*...*`` italic wrapper the response uses for every quotation is
    stripped here and the italic is applied to the whole run instead.
    """
    lines = open(path, encoding="utf-8").read().split("\n")
    blocks, i, n = [], 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()

        if not s:
            i += 1
            continue

        if re.fullmatch(r"-{3,}", s):
            blocks.append(("hr", None))
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append(("h", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue

        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                raw = lines[i].strip()
                if not SEP_RE.match(raw):
                    rows.append(split_row(raw))
                i += 1
            if rows:
                ncols = max(len(r) for r in rows)
                rows = [r + [""] * (ncols - len(r)) for r in rows]
                blocks.append(("table", rows))
            continue

        if s.startswith(">"):
            paras = []
            while i < n and lines[i].strip().startswith(">"):
                body = re.sub(r"^\s*>\s?", "", lines[i]).rstrip()
                if body.strip():
                    paras.append(strip_italic_wrapper(body.strip()))
                i += 1
            if paras:
                blocks.append(("quote", paras))
            continue

        m = NUM_ITEM_RE.match(s)
        if m:
            blocks.append(("num", (m.group(1) + ".", m.group(2).strip())))
            i += 1
            continue

        if BULLET_RE.match(line):
            buf = [BULLET_RE.sub("", line).strip()]
            i += 1
            while i < n:
                nxt = lines[i]
                if not nxt.strip():
                    break
                if (BULLET_RE.match(nxt) or NUM_ITEM_RE.match(nxt.strip())
                        or nxt.strip().startswith(("|", ">", "#"))):
                    break
                buf.append(nxt.strip())
                i += 1
            blocks.append(("bullet", " ".join(buf)))
            continue

        buf = [s]
        i += 1
        while i < n:
            nxt = lines[i]
            if not nxt.strip():
                break
            if (nxt.strip().startswith(("|", ">", "#")) or BULLET_RE.match(nxt)
                    or NUM_ITEM_RE.match(nxt.strip())
                    or re.fullmatch(r"-{3,}", nxt.strip())):
                break
            buf.append(nxt.strip())
            i += 1
        blocks.append(("p", " ".join(buf)))
    return blocks


def strip_italic_wrapper(text: str) -> str:
    """Remove the single outer ``*...*`` the response wraps every quotation in.

    Only the outermost pair is removed, and only when both ends carry it, so a
    quotation that itself starts with the editor's ``* `` bullet keeps it.
    """
    if len(text) >= 2 and text.startswith("*") and text.endswith("*"):
        return text[1:-1]
    return text


# ==========================================================================
# inline markdown -> runs
# ==========================================================================

EMDASH_RE = re.compile(r"--")


def parse_inline(text: str) -> list:
    """Return ``[(text, bold, italic, code), ...]``.

    Well-formed markdown is assumed; an unbalanced marker degrades to a literal
    character rather than raising.
    """
    runs, buf = [], []
    bold = italic = False
    i, n = 0, len(text)

    def flush(code=False):
        if buf:
            t = "".join(buf)
            if not code:
                t = EMDASH_RE.sub("\u2013", t)
            runs.append((t, bold, italic, code))
            buf.clear()

    while i < n:
        ch = text[i]
        if ch == "`":
            j = text.find("`", i + 1)
            if j == -1:
                buf.append(ch)
                i += 1
                continue
            flush()
            code_txt = text[i + 1:j]
            if code_txt:
                runs.append((code_txt, bold, italic, True))
            i = j + 1
            continue
        if text.startswith("**", i):
            flush()
            bold = not bold
            i += 2
            continue
        if ch == "*":
            flush()
            italic = not italic
            i += 1
            continue
        buf.append(ch)
        i += 1
    flush()
    return [r for r in runs if r[0]]


def plain_text(text: str) -> str:
    return "".join(r[0] for r in parse_inline(text))


# ==========================================================================
# docx helpers
# ==========================================================================

CODE_FONT = "Consolas"
ZWSP = "\u200b"
BREAK_AFTER = set("/_-,.:;{}()=|\\&+#?")
LONG_TOKEN = 14

_OPENERS = "([{<\u2018\u201c\"'"
_CLOSERS = ")]}>\u2019\u201d\"'.,;:!?"
_URL_PREFIXES = ("http://", "https://", "www.", "doi:")
_DOI_RE = re.compile(r"10\.\d{4,9}/")
_HYPHEN_WORD_RE = re.compile(r"[A-Za-z0-9]+(?:-[A-Za-z0-9]+)+")


def _is_link(tok: str) -> bool:
    core = tok.lstrip(_OPENERS).lower()
    return core.startswith(_URL_PREFIXES) or _DOI_RE.match(core) is not None


def _is_hyphenated_word(tok: str) -> bool:
    """``cluster-bootstrap``, ``MAGICC-minus-CheckM2``.

    Word and LibreOffice both break a line after a hyphen natively, so the
    helper would only add a hidden second copy of a break opportunity the text
    already carries.
    """
    return _HYPHEN_WORD_RE.fullmatch(tok.strip(_OPENERS + _CLOSERS)) is not None


def soft_break(text: str, break_links: bool = False) -> str:
    """Insert zero-width spaces inside over-long tokens.

    A zero-width space is an invisible break opportunity that both Word and
    LibreOffice honour; it does not change the visible text.

    Body prose keeps the usual exemption for web addresses -- the longest link
    outside the quotations is 29 characters and none of them needs help
    wrapping, and a hidden character inside a link the reader is meant to follow
    would stop it resolving once it is copied out of the PDF.  Inside the
    quotations the exemption is lifted (``break_links=True``): the editor's
    letter carries policy URLs up to 108 characters, which are 7.6 in wide at
    10.5 pt and would otherwise be painted straight off the right edge of the
    page.  Those URLs are quoted for the record, not for the reader to copy.
    """
    out = []
    for tok in re.split(r"(\s+)", text):
        if (len(tok) > LONG_TOKEN and not tok.isspace()
                and not (_is_link(tok) and not break_links)
                and not _is_hyphenated_word(tok)):
            buf, run_len = [], 0
            for ch in tok:
                buf.append(ch)
                run_len += 1
                if ch in BREAK_AFTER and run_len >= 6:
                    buf.append(ZWSP)
                    run_len = 0
            tok = "".join(buf)
        out.append(tok)
    return "".join(out)


def set_run_font(run, name):
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)
    rf.attrib.pop(qn("w:hint"), None)


def fill_runs(par, text, size_pt=None, force_bold=False, force_italic=False):
    """Replace the paragraph's runs with formatted runs parsed from markdown."""
    for r in list(par._p.findall(qn("w:r"))):
        par._p.remove(r)
    for hl in list(par._p.findall(qn("w:hyperlink"))):
        par._p.remove(hl)
    for txt, bold, italic, code in parse_inline(text):
        run = par.add_run(soft_break(txt))
        run.bold = True if (bold or force_bold) else None
        run.italic = True if (italic or force_italic) else None
        if code:
            set_run_font(run, CODE_FONT)
        if size_pt is not None:
            run.font.size = Pt(size_pt)
    return par


def new_paragraph(doc, style=None):
    p = OxmlElement("w:p")
    par = Paragraph(p, doc._body)
    if style is not None:
        par.style = doc.styles[style]
    return par


def set_indent(par, left_twips=0, hanging_twips=0, right_twips=0):
    ppr = par._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:left"), str(left_twips))
    if right_twips:
        ind.set(qn("w:right"), str(right_twips))
    if hanging_twips:
        ind.set(qn("w:hanging"), str(hanging_twips))


def set_spacing(par, before=None, after=None):
    ppr = par._p.get_or_add_pPr()
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        ppr.insert(0, sp)
    if before is not None:
        sp.set(qn("w:before"), str(before))
    if after is not None:
        sp.set(qn("w:after"), str(after))


def keep_with_next(par):
    ppr = par._p.get_or_add_pPr()
    if ppr.find(qn("w:keepNext")) is None:
        ppr.append(OxmlElement("w:keepNext"))


def shade(par, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    par._p.get_or_add_pPr().append(shd)


def add_left_bar(par, color=QUOTE_BAR, sz=18):
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    bdr.append(left)
    ppr.append(bdr)


def add_page_break(doc):
    par = new_paragraph(doc, "Normal")
    r = par.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r._r.append(br)
    return par


def add_rule(doc):
    """A thin horizontal rule: an empty paragraph with a bottom border."""
    par = new_paragraph(doc, "Normal")
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFBFBF")
    bdr.append(bottom)
    ppr.append(bdr)
    set_spacing(par, before=120, after=120)
    return par


# ==========================================================================
# text metrics -- Arial's advance widths are the Helvetica ones
# ==========================================================================

_HELV = {
    " ": 278, "!": 278, '"': 355, "#": 556, "$": 556, "%": 889, "&": 667,
    "'": 191, "(": 333, ")": 333, "*": 389, "+": 584, ",": 278, "-": 333,
    ".": 278, "/": 278, ":": 278, ";": 278, "<": 584, "=": 584, ">": 584,
    "?": 556, "@": 1015, "[": 278, "\\": 278, "]": 278, "^": 469, "_": 556,
    "`": 333, "{": 334, "|": 260, "}": 334, "~": 584,
    "A": 667, "B": 667, "C": 722, "D": 722, "E": 667, "F": 611, "G": 778,
    "H": 722, "I": 278, "J": 500, "K": 667, "L": 556, "M": 833, "N": 722,
    "O": 778, "P": 667, "Q": 778, "R": 722, "S": 667, "T": 611, "U": 722,
    "V": 667, "W": 944, "X": 667, "Y": 667, "Z": 611,
    "a": 556, "b": 556, "c": 500, "d": 556, "e": 556, "f": 278, "g": 556,
    "h": 556, "i": 222, "j": 222, "k": 500, "l": 222, "m": 833, "n": 556,
    "o": 556, "p": 556, "q": 556, "r": 333, "s": 500, "t": 278, "u": 556,
    "v": 500, "w": 722, "x": 500, "y": 500, "z": 500,
}
for _d in "0123456789":
    _HELV[_d] = 556
_DEFAULT_ADV = 600
_BOLD_FACTOR = 1.07
_MONO_ADV = 550


def _adv(ch, bold, code):
    w = _MONO_ADV if code else _HELV.get(ch, _DEFAULT_ADV)
    return w * (_BOLD_FACTOR if bold else 1.0) / 1000.0


def md_metrics(cell_md, header=False):
    """(total width, longest unbreakable width) of a markdown fragment, in em."""
    stream = []
    for txt, bold, italic, code in parse_inline(cell_md):
        for ch in soft_break(txt):
            stream.append((ch, _adv(ch, bold or header, code)))
    total = sum(a for _, a in stream)
    longest = run = 0.0
    for ch, a in stream:
        if ch.isspace() or ch == ZWSP:
            run = 0.0
        else:
            run += a
            longest = max(longest, run)
    return total, longest


def text_em(text, bold=False):
    return sum(_adv(ch, bold, False) for ch in text)


def char_width(cell_text: str) -> int:
    return len(plain_text(cell_text))


CELL_PAD = 90
MIN_COL_TWIPS = 480
TARGET_LINE_EM = 17.0


def column_metrics(rows, size_pt):
    ncols = len(rows[0])
    twips_per_em = size_pt * 20.0
    need, want = [], []
    for c in range(ncols):
        totals, longests = [], []
        for ri, r in enumerate(rows):
            t, l = md_metrics(r[c], header=(ri == 0))
            totals.append(t)
            longests.append(l)
        need.append(max(max(longests) * twips_per_em + CELL_PAD, MIN_COL_TWIPS))
        want.append(min(max(totals), TARGET_LINE_EM) * twips_per_em + CELL_PAD)
    return need, want


def allocate(total, need, want):
    """Split ``total`` proportionally to ``want`` but never below ``need``."""
    n = len(need)
    if sum(need) >= total:
        scale = total / sum(need)
        w = [x * scale for x in need]
    else:
        w = [None] * n
        free = set(range(n))
        remaining = float(total)
        while free:
            tot_w = sum(want[i] for i in free) or 1.0
            pinned = [i for i in free if remaining * want[i] / tot_w < need[i]]
            if not pinned:
                for i in free:
                    w[i] = remaining * want[i] / tot_w
                break
            for i in pinned:
                w[i] = need[i]
                remaining -= need[i]
                free.discard(i)
    out = [int(round(x)) for x in w]
    k = max(range(n), key=lambda i: out[i])
    out[k] += total - sum(out)
    return out


def table_layout(rows, total=TABLE_TOTAL_TWIPS):
    """Choose the body font size and the column widths together.

    Wide tables shrink rather than overflow; the size drops further only if the
    unbreakable minimum widths still do not fit.
    """
    ncols = len(rows[0])
    max_row_em = max(sum(md_metrics(c, header=(ri == 0))[0] for c in r)
                     for ri, r in enumerate(rows))
    if ncols >= 6:
        start = 8.0
    elif ncols == 5:
        start = 9.0
    else:
        start = 10.0
    if max_row_em > 145 and start > 8.0:
        start -= 1.0
    if max_row_em > 235:
        start = 8.0

    for size_pt in [x for x in (10.0, 9.5, 9.0, 8.0, 7.0, 6.5) if x <= start]:
        need, want = column_metrics(rows, size_pt)
        if sum(need) <= total:
            return size_pt, allocate(total, need, want)
    size_pt = 6.5
    need, want = column_metrics(rows, size_pt)
    scale = total / sum(need)
    need = [x * scale for x in need]
    return size_pt, allocate(total, need, want)


def build_table(doc, tbl_el, rows):
    """Rebuild ``tbl_el`` (a clone of the v4 table) from ``rows`` in place.

    Row and cell formatting -- borders, cell margins, the shaded header fill --
    come from the element's own first (header) and second (body) rows, so every
    table inherits the v4 look.
    """
    trs = tbl_el.findall(qn("w:tr"))
    hdr_tpl = copy.deepcopy(trs[0])
    body_tpl = copy.deepcopy(trs[1]) if len(trs) > 1 else copy.deepcopy(trs[0])
    for tr in trs:
        tbl_el.remove(tr)

    def strip_to_one_cell(tr_tpl):
        tcs = tr_tpl.findall(qn("w:tc"))
        keep = copy.deepcopy(tcs[0])
        for tc in tr_tpl.findall(qn("w:tc")):
            tr_tpl.remove(tc)
        return tr_tpl, keep

    hdr_row_tpl, hdr_cell_tpl = strip_to_one_cell(hdr_tpl)
    body_row_tpl, body_cell_tpl = strip_to_one_cell(body_tpl)

    ncols = len(rows[0])
    size_pt, widths = table_layout(rows)

    aligns = []
    for c in range(ncols):
        longest = max(char_width(r[c]) for r in rows)
        aligns.append("center" if longest <= 20 else "left")

    tblpr = tbl_el.find(qn("w:tblPr"))
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.insert(0, tblw)
    tblw.set(qn("w:w"), str(TABLE_TOTAL_TWIPS))
    tblw.set(qn("w:type"), "dxa")
    for old in tblpr.findall(qn("w:tblLayout")):
        tblpr.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    look = tblpr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(layout)
    else:
        tblpr.append(layout)

    grid = tbl_el.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tblpr.addnext(grid)
    for gc in list(grid.findall(qn("w:gridCol"))):
        grid.remove(gc)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)

    for ri, row in enumerate(rows):
        is_hdr = ri == 0
        tr = copy.deepcopy(hdr_row_tpl if is_hdr else body_row_tpl)
        cell_tpl = hdr_cell_tpl if is_hdr else body_cell_tpl
        for ci, cell_md in enumerate(row):
            tc = copy.deepcopy(cell_tpl)
            tcpr = tc.find(qn("w:tcPr"))
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.insert(0, tcw)
            tcw.set(qn("w:w"), str(widths[ci]))
            tcw.set(qn("w:type"), "dxa")
            ps = tc.findall(qn("w:p"))
            for extra in ps[1:]:
                tc.remove(extra)
            par = Paragraph(ps[0], doc._body)
            fill_runs(par, cell_md, size_pt=size_pt, force_bold=is_hdr)
            ppr = par._p.get_or_add_pPr()
            for j in list(ppr.findall(qn("w:jc"))):
                ppr.remove(j)
            jc = OxmlElement("w:jc")
            jc.set(qn("w:val"), aligns[ci] if not is_hdr else "center")
            ppr.append(jc)
            tr.append(tc)
        tbl_el.append(tr)
    return tbl_el


# ==========================================================================
# build
# ==========================================================================

def tune_styles(doc):
    """Keep the v4 styles but stop headings being stranded at a page foot."""
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[name]
        st.paragraph_format.keep_with_next = True


def drop_images(doc):
    """Discard the v4 manuscript's four figures.

    Only the furniture is wanted from the v4 file.  Dropping the relationships
    takes the image parts out of the saved package, which keeps the response a
    small text document instead of a 2 MB one.
    """
    dropped = 0
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/image"):
            doc.part.drop_rel(rid)
            dropped += 1
    return dropped


def build():
    global USABLE_TWIPS

    shutil.copyfile(V4_DOCX, OUT)
    doc = docx.Document(OUT)

    sect = doc.sections[0]
    USABLE_TWIPS = int(
        (sect.page_width - sect.left_margin - sect.right_margin) / 635
    )  # EMU -> twips
    assert TABLE_TOTAL_TWIPS <= USABLE_TWIPS, (
        f"table width {TABLE_TOTAL_TWIPS} exceeds usable {USABLE_TWIPS} twips"
    )

    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))

    # --- inventory the v4 body, then keep only the furniture --------------
    children = [c for c in body.iterchildren() if c is not sectPr]
    v4_paras = [c for c in children if c.tag == qn("w:p")]
    v4_tables = [c for c in children if c.tag == qn("w:tbl")]
    v4_images = [p for p in v4_paras if p.findall(".//" + qn("a:blip"))]
    before = dict(paragraphs=len(v4_paras), tables=len(v4_tables),
                  images=len(v4_images))
    assert v4_tables, "the v4 file must supply at least one table template"

    table_template = copy.deepcopy(v4_tables[0])

    for c in children:
        body.remove(c)
    n_dropped = drop_images(doc)
    tune_styles(doc)

    def append(el):
        sectPr.addprevious(el)

    # --- markdown ---------------------------------------------------------
    blocks = parse_markdown(MD)

    # A thematic break starts a new page once the section it closes carries
    # subsections of its own; before that -- the title block and the one-
    # paragraph revision note -- it is set as a rule, so the opening pages are
    # not left nearly empty.  On this document that puts Reviewer 1, Reviewer 2,
    # the editorial requirements and the closing section each on a fresh page,
    # which is how the previous build laid it out.
    h3_since_break = 0
    md_table_count = 0
    n_quote_blocks = 0
    n_quote_paras = 0

    for kind, payload in blocks:
        if kind == "hr":
            if h3_since_break:
                append(add_page_break(doc)._p)
                h3_since_break = 0
            else:
                append(add_rule(doc)._p)
            continue

        if kind == "h":
            level, text = payload
            style = {1: "Title", 2: "Heading 1", 3: "Heading 2"}.get(
                level, "Heading 3")
            par = new_paragraph(doc, style)
            fill_runs(par, text)
            append(par._p)
            if level == 3:
                h3_since_break += 1
            continue

        if kind == "p":
            text = payload
            par = new_paragraph(doc, "Normal")
            fill_runs(par, text)
            set_spacing(par, after=120)
            if plain_text(text).strip() in LABELS:
                set_spacing(par, before=120, after=60)
                keep_with_next(par)
            append(par._p)
            continue

        if kind == "bullet":
            par = new_paragraph(doc, "Normal")
            fill_runs(par, "\u2022 " + payload)
            set_indent(par, left_twips=360, hanging_twips=200)
            set_spacing(par, after=60)
            append(par._p)
            continue

        if kind == "num":
            label, text = payload
            par = new_paragraph(doc, "Normal")
            fill_runs(par, label + " " + text)
            # the number sits where a bullet would (160 twips) and the text
            # hangs clear of a two-digit label
            set_indent(par, left_twips=430, hanging_twips=270)
            set_spacing(par, after=60)
            append(par._p)
            continue

        if kind == "quote":
            n_quote_blocks += 1
            last = len(payload) - 1
            for k, qtext in enumerate(payload):
                n_quote_paras += 1
                par = new_paragraph(doc, "Normal")
                run = par.add_run(soft_break(qtext, break_links=True))
                run.italic = True
                run.font.size = Pt(QUOTE_PT)
                run.font.color.rgb = QUOTE_COLOR
                set_indent(par, left_twips=QUOTE_LEFT_TWIPS,
                           right_twips=QUOTE_RIGHT_TWIPS)
                set_spacing(par, before=120 if k == 0 else 40,
                            after=120 if k == last else 40)
                shade(par, QUOTE_FILL)
                add_left_bar(par)
                append(par._p)
            continue

        if kind == "table":
            md_table_count += 1
            tbl_el = copy.deepcopy(table_template)
            build_table(doc, tbl_el, payload)
            append(tbl_el)
            append(new_paragraph(doc, "Normal")._p)
            continue

        raise AssertionError(f"unhandled block {kind}")

    doc.save(OUT)

    d2 = docx.Document(OUT)
    after = dict(paragraphs=len(d2.paragraphs), tables=len(d2.tables),
                 images=len(d2.inline_shapes))
    return dict(before=before, after=after, tables=md_table_count,
                quote_blocks=n_quote_blocks, quote_paras=n_quote_paras,
                images_dropped=n_dropped)


# ==========================================================================
# verification
# ==========================================================================

NUM_RE = re.compile(r"[-\u2212+]?\d[\d,]*\.?\d*")


def norm_num(s):
    return s.replace(",", "").replace("\u2212", "-").rstrip(".")


def squash(s):
    return " ".join(s.replace(ZWSP, "").split())


def expected_stream(blocks):
    """The text the document must carry, block by block, in order."""
    out = []
    for kind, payload in blocks:
        if kind == "hr":
            continue
        if kind == "h":
            out.append(("h%d" % payload[0], plain_text(payload[1])))
        elif kind == "p":
            out.append(("p", plain_text(payload)))
        elif kind == "bullet":
            out.append(("p", "\u2022 " + plain_text(payload)))
        elif kind == "num":
            out.append(("p", payload[0] + " " + plain_text(payload[1])))
        elif kind == "quote":
            for q in payload:
                out.append(("q", q))
        elif kind == "table":
            for row in payload:
                for cell in row:
                    out.append(("c", plain_text(cell)))
    return out


def actual_stream(doc):
    """The text the document actually carries, in document order."""
    out = []
    style_of = {"Title": "h1", "Heading 1": "h2", "Heading 2": "h3",
                "Heading 3": "h4"}
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            from docx.table import Table
            t = Table(child, doc)
            for row in t.rows:
                for cell in row.cells:
                    out.append(("c", cell.text))
            continue
        if child.tag != qn("w:p"):
            continue
        p = Paragraph(child, doc)
        if not p.text.strip():
            continue
        ppr = child.find(qn("w:pPr"))
        is_quote = ppr is not None and ppr.find(qn("w:pBdr")) is not None \
            and ppr.find(qn("w:pBdr")).find(qn("w:left")) is not None
        kind = style_of.get(p.style.name, "q" if is_quote else "p")
        out.append((kind, p.text))
    return out


def verify(stats):
    print("=" * 74)
    print("VERIFICATION")
    print("=" * 74)

    doc = docx.Document(OUT)
    print(f"opened OK: {OUT}")
    print(f"  file size: {os.path.getsize(OUT) / 1024:.1f} KB")
    print(f"  v4 furniture discarded: paragraphs={stats['before']['paragraphs']}, "
          f"tables={stats['before']['tables']}, "
          f"images={stats['before']['images']} "
          f"({stats['images_dropped']} image relationships dropped)")
    print(f"  rendered: paragraphs={stats['after']['paragraphs']}, "
          f"tables={stats['after']['tables']}, "
          f"inline shapes={stats['after']['images']}")
    assert stats["after"]["images"] == 0
    assert stats["after"]["tables"] == stats["tables"], (
        stats["after"]["tables"], stats["tables"])
    print(f"  OK  {stats['after']['tables']} Word tables == "
          f"{stats['tables']} markdown tables; no images")

    blocks = parse_markdown(MD)

    # ---- whole-document text diff, block by block ------------------------
    want = expected_stream(blocks)
    got = actual_stream(doc)
    diffs = []
    n = max(len(want), len(got))
    for i in range(n):
        w = want[i] if i < len(want) else ("--", "<missing>")
        g = got[i] if i < len(got) else ("--", "<missing>")
        if w[0] != g[0] or squash(w[1]) != squash(g[1]):
            diffs.append((i, w, g))
    print(f"  text stream: markdown {len(want)} blocks, docx {len(got)} blocks; "
          f"differences: {len(diffs)}")
    for i, w, g in diffs[:8]:
        print(f"    [{i}] md({w[0]}) {w[1][:70]!r}")
        print(f"         docx({g[0]}) {g[1][:70]!r}")
    assert not diffs, "the .docx does not agree with the markdown"

    # ---- headings --------------------------------------------------------
    md_h = {}
    for k, p in blocks:
        if k == "h":
            md_h[p[0]] = md_h.get(p[0], 0) + 1
    doc_h = {}
    for kind, _ in got:
        if kind.startswith("h"):
            doc_h[kind] = doc_h.get(kind, 0) + 1
    print(f"  headings: markdown {dict(sorted(md_h.items()))}  "
          f"docx {dict(sorted(doc_h.items()))}")
    assert md_h.get(3) == N_H3, md_h.get(3)
    assert doc_h.get("h3") == N_H3, doc_h.get("h3")
    print(f"  OK  all {N_H3} '###' headings present as Heading 2")

    # ---- entry blocks ----------------------------------------------------
    doc_paras = [t for k, t in got if k == "p"]
    n_resp = sum(1 for t in doc_paras if squash(t) == "Response.")
    n_chg = sum(1 for t in doc_paras if squash(t) == "Changes made.")
    print(f"  entry blocks: quote blocks {stats['quote_blocks']}, "
          f"'Response.' {n_resp}, 'Changes made.' {n_chg}")
    assert stats["quote_blocks"] == n_resp == n_chg == N_ENTRIES, (
        stats["quote_blocks"], n_resp, n_chg)
    print(f"  OK  all {N_ENTRIES} entry blocks complete "
          f"(comment + Response. + Changes made.)")

    # ---- the 72 verbatim quote lines, character by character -------------
    md_lines = [l for l in open(MD, encoding="utf-8").read().split("\n")
                if l.startswith(">")]
    md_quotes = []
    for l in md_lines:
        b = re.sub(r"^\s*>\s?", "", l).strip()
        if b:
            md_quotes.append(strip_italic_wrapper(b))
    doc_quotes = [t for k, t in got if k == "q"]
    assert len(md_lines) == N_QUOTE_LINES, len(md_lines)
    assert len(md_quotes) == N_QUOTE_PARAS, len(md_quotes)
    bad = [(i, a, b) for i, (a, b) in enumerate(zip(md_quotes, doc_quotes))
           if a != b.replace(ZWSP, "")]
    print(f"  verbatim quote lines: markdown {len(md_lines)} '>' lines "
          f"({len(md_lines) - len(md_quotes)} of them blank paragraph "
          f"separators, {len(md_quotes)} carrying text), "
          f"docx {len(doc_quotes)}; character-level differences: {len(bad)}")
    for i, a, b in bad[:5]:
        print(f"    [{i}] md   {a[:70]!r}")
        print(f"         docx {b[:70]!r}")
    assert len(doc_quotes) == N_QUOTE_PARAS and not bad
    zw = sum(1 for q in doc_quotes if ZWSP in q)
    print(f"  OK  all {N_QUOTE_LINES} quote lines reproduced and every one of "
          f"the {N_QUOTE_PARAS} text-bearing lines is character-identical "
          f"(after removing the invisible break opportunities added to "
          f"{zw} of them)")

    # ---- numbers ---------------------------------------------------------
    md_nums = [norm_num(x) for _, t in want for x in NUM_RE.findall(t)]
    doc_nums = [norm_num(x) for _, t in got
                for x in NUM_RE.findall(t.replace(ZWSP, ""))]
    print(f"  numeric tokens: markdown {len(md_nums)}, docx {len(doc_nums)}; "
          f"identical: {md_nums == doc_nums}")
    assert md_nums == doc_nums

    # ---- bold-run audit --------------------------------------------------
    md_bold = set()
    for k, p in blocks:
        if k in ("p", "bullet"):
            src = p if k != "num" else p[1]
            for txt, b, i2, c in parse_inline(src):
                if b:
                    md_bold.add(squash(txt))
        elif k == "num":
            for txt, b, i2, c in parse_inline(p[1]):
                if b:
                    md_bold.add(squash(txt))
        elif k == "table":
            for row in p:
                for cell in row:
                    for txt, b, i2, c in parse_inline(cell):
                        if b:
                            md_bold.add(squash(txt))
    hdr_bold = set()
    for k, p in blocks:
        if k == "table":
            for cell in p[0]:
                hdr_bold.add(squash(plain_text(cell)))

    body_bold, table_bold, stray = [], [], []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            for r in p.runs:
                if r.bold:
                    body_bold.append(squash(r.text))
        elif child.tag == qn("w:tbl"):
            from docx.table import Table
            t = Table(child, doc)
            for ri, row in enumerate(t.rows):
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for r in para.runs:
                            if r.bold:
                                table_bold.append((ri, squash(r.text)))
    allowed = md_bold | hdr_bold
    for txt in body_bold:
        if txt not in allowed:
            stray.append(("body", txt))
    for ri, txt in table_bold:
        if ri != 0 and txt not in allowed:
            stray.append(("table", txt))
    n_hdr_cells = sum(1 for ri, _ in table_bold if ri == 0)
    print(f"  bold runs: {len(body_bold)} in body paragraphs "
          f"+ {len(table_bold)} in tables "
          f"({n_hdr_cells} header cells, "
          f"{len(table_bold) - n_hdr_cells} comment identifiers) "
          f"= {len(body_bold) + len(table_bold)} total")
    labs = sum(1 for t in body_bold if t in LABELS)
    print(f"    of the body runs, {labs} are the 'Response.' / 'Changes made.' "
          f"labels and {len(body_bold) - labs} are structural "
          f"({sorted(set(body_bold) - set(LABELS))})")
    print(f"    bold spans the markdown asks for: {len(md_bold)} distinct; "
          f"runs bold in the docx but not in the markdown: {len(stray)}")
    for s in stray[:8]:
        print("    STRAY BOLD:", s)
    assert not stray, "decorative bold crept back in"

    # ---- no raw markdown left over --------------------------------------
    scan = [t for _, t in got]
    leftovers = {}
    for pat, name in ((r"\*\*", "**"),
                      (r"^\s*\|.*\|\s*$", "| table row"),
                      (r"\|\s*-{3,}\s*\|", "|--- rule"),
                      (r"^#{1,6}\s", "#"), (r"\bTODO\b", "TODO"),
                      (r"\bTBD\b", "TBD"), (r"PLACEHOLDER", "PLACEHOLDER")):
        hits = [s for s in scan if re.search(pat, s, re.M)]
        if hits:
            leftovers[name] = hits[:3]
    print(f"  raw-markdown / TODO leftovers: "
          f"{ {k: len(v) for k, v in leftovers.items()} or 'none'}")
    assert not leftovers, leftovers

    # ---- standing conventions -------------------------------------------
    joined = "\n".join(scan)
    conv = []
    for ln in joined.split("\n"):
        if "0.442" in ln or "44.2 %" in ln:
            if not any(k in ln for k in ("0.020", "2.0 %", "0.769",
                                         "balanced accuracy")):
                conv.append("false-fail rate quoted bare")
    if "cold_vs_warm" in joined:
        conv.append("cold_vs_warm.tsv cited")
    for tok in ("1,451 genomes", "1,700\u00d7", "2,100\u00d7", "magicc-genome",
                "8.18 GB", "77.3 %", "315 rejected"):
        for ln in joined.split("\n"):
            if tok in ln and not any(k in ln.lower() for k in (
                    "withdraw", "never measured", "does not exist", "not exist",
                    "superseded", "correct", "error", "reconcil", "no longer",
                    "inconsisten", "does not support")):
                conv.append(f"withdrawn token {tok!r} used without a "
                            f"withdrawal marker")
    print(f"  standing-convention checks: "
          f"{'all clear' if not conv else conv}")
    assert not conv, conv

    # ---- table widths ----------------------------------------------------
    over = []
    for i, t in enumerate(doc.tables):
        g = t._tbl.find(qn("w:tblGrid"))
        w = sum(int(c.get(qn("w:w"))) for c in g.findall(qn("w:gridCol")))
        if w > USABLE_TWIPS:
            over.append((i, w))
    print(f"  usable text width: {USABLE_TWIPS} twips "
          f"({USABLE_TWIPS / 1440:.3f} in); tables wider than that: {len(over)}")
    assert not over, over

    # ---- nothing wider than its container --------------------------------
    body_over = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=11.0)
        limit = USABLE_TWIPS
        ppr = p._p.find(qn("w:pPr"))
        if ppr is not None:
            i2 = ppr.find(qn("w:ind"))
            if i2 is not None:
                limit -= int(i2.get(qn("w:left")) or 0)
                limit -= int(i2.get(qn("w:right")) or 0)
        bold = any(r.bold for r in p.runs)
        worst = max((text_em(t, bold) for t in
                     re.split(r"[\s\u200b]+", p.text) if t), default=0.0)
        if worst * size * 20.0 > limit:
            body_over.append((round(worst, 1), p.text[:70]))
    print(f"  body paragraphs with an unbreakable run wider than the text "
          f"column: {len(body_over)}")
    for w, t in body_over[:5]:
        print(f"    {w} em: {t!r}")
    assert not body_over

    cell_over = []
    for ti, t in enumerate(doc.tables):
        g = t._tbl.find(qn("w:tblGrid"))
        cw = [int(c.get(qn("w:w"))) for c in g.findall(qn("w:gridCol"))]
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                if ci >= len(cw):
                    continue
                size = max((r.font.size.pt for p in cell.paragraphs
                            for r in p.runs if r.font.size), default=10.0)
                bold = ri == 0 or any(r.bold for p in cell.paragraphs
                                      for r in p.runs)
                worst = max((text_em(x, bold) for x in
                             re.split(r"[\s\u200b]+", cell.text) if x),
                            default=0.0)
                if worst * size * 20.0 > cw[ci] - CELL_PAD + 20:
                    cell_over.append((ti, ri, ci, round(worst, 1),
                                      cell.text[:40]))
    print(f"  table cells with an unbreakable run wider than their column: "
          f"{len(cell_over)}")
    for x in cell_over[:5]:
        print("   ", x)
    assert not cell_over

    sizes = sorted({round(r.font.size.pt, 1) for t in doc.tables
                    for row in t.rows for c in row.cells
                    for p in c.paragraphs for r in p.runs if r.font.size})
    print(f"  table body font sizes in use: {sizes} pt")

    # ---- PDF -------------------------------------------------------------
    tmp = tempfile.mkdtemp(prefix="respdocx_")
    r = subprocess.run(["/usr/bin/libreoffice", "--headless", "--convert-to",
                        "pdf", "--outdir", tmp, OUT],
                       capture_output=True, text=True, timeout=1800)
    pdf = os.path.join(tmp, os.path.basename(OUT).replace(".docx", ".pdf"))
    if not os.path.exists(pdf):
        print("  PDF CONVERSION FAILED:", r.stdout, r.stderr)
        raise SystemExit(1)
    info = subprocess.run(["pdfinfo", pdf], capture_output=True,
                          text=True).stdout
    pages = int(re.search(r"^Pages:\s+(\d+)", info, re.M).group(1))
    ps = re.search(r"^Page size:\s+(.*)$", info, re.M).group(1)
    print(f"  PDF rendered OK: {os.path.getsize(pdf) / 1e6:.2f} MB, "
          f"{pages} pages, page size {ps}")
    dims = re.match(r"([\d.]+) x ([\d.]+) pts", ps)
    assert dims and abs(float(dims.group(1)) - 595.3) < 0.1 \
        and abs(float(dims.group(2)) - 841.9) < 0.1, ps
    print(f"  page size confirmed A4: {dims.group(1)} x {dims.group(2)} pts")

    bbox = subprocess.run(["pdftotext", "-bbox", pdf, "-"],
                          capture_output=True, text=True).stdout
    right_limit = 595.304 - 72.0 + 1.5
    page_no, worst, offenders = 0, 0.0, []
    for line in bbox.split("\n"):
        if "<page " in line:
            page_no += 1
        m = re.search(r'<word xMin="([\d.]+)" yMin="([\d.]+)" '
                      r'xMax="([\d.]+)" yMax="([\d.]+)">(.*)</word>', line)
        if not m:
            continue
        xmax = float(m.group(3))
        if xmax > right_limit:
            worst = max(worst, xmax)
            offenders.append((page_no, round(xmax, 1), m.group(5)[:40]))
    print(f"  words painted past the right margin "
          f"({right_limit - 1.5:.0f} pt): {len(offenders)}"
          + (f"   worst xMax {worst:.1f} pt" if offenders else ""))
    for o in offenders[:8]:
        print(f"    page {o[0]}  xMax {o[1]}  {o[2]!r}")
    assert not offenders, "content overflows the page width"
    shutil.rmtree(tmp, ignore_errors=True)

    print("=" * 74)
    print("ALL CHECKS PASSED")
    return pages


def main():
    md_before = open(MD, "rb").read()
    stats = build()
    md_after = open(MD, "rb").read()
    assert md_before == md_after, (
        "the markdown changed during the build -- it is the source of truth "
        "and this script must never write it")
    print(f"source (unmodified): {MD}  "
          f"({len(md_before) / 1024:.1f} KB)")
    print(f"wrote: {OUT}  ({os.path.getsize(OUT) / 1024:.1f} KB)")
    print()
    pages = verify(stats)
    print(f"{OUT}: {pages} pages, A4.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
