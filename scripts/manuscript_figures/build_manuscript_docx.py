#!/usr/bin/env python
"""Build ``manuscript_revised.docx`` (+ a change-marked twin) by rendering
``manuscript_revised.md`` into the document furniture of the submitted v4 file.

Why this is a renderer and no longer a patcher
----------------------------------------------
The first revision was a set of insertions, so the previous version of this
script copied ``1st_submission/manuscript_v4.docx`` and applied 177 entries from
``manuscript_change_map.md`` to it in place.  That is no longer expressible: the
manuscript has been compressed from 6,945 to 4,947 main-text words, de-bolded
throughout, cut from 7 figures + 2 tables to 5 figures + 1 table, and its legends
now sit inline under their display items.  A change map cannot say that.

So the markdown is now the source of truth and is rendered directly.  What the
v4 file still supplies is the *document design*: it is copied to the output path
and its body is emptied, which keeps styles.xml (Title / Heading 1-2 / Normal /
Bibliography, Arial 11 pt), the A4 ``sectPr`` with 1 in margins, the header and
footer parts, theme, fontTable, numbering and settings exactly as a reader saw
them the first time.  Nothing is regenerated from a python-docx default
template, which would be US Letter and Calibri.

What the script does, in order:

1.  Copy the pristine v4 file to the output path (v4 is never opened for
    writing) and capture, before anything is touched, the three pieces of
    furniture that have to be cloned rather than invented: an image paragraph
    (so a new figure carries the drawing attributes Word wrote), the single
    ``w:tbl`` (so Table 1 keeps the v4 borders, shading and cell margins), and
    the v4 wording of every paragraph and table cell (so the marked copy can
    diff against it).
2.  Empty the body, keeping ``sectPr``.
3.  Render ``manuscript_revised.md`` block by block.  ``[[FIGURE:n]]`` embeds
    ``figures/Figure_n.png`` at 400 dpi, sized to the text width, and the
    paragraph beginning ``**Figure n.**`` is emitted immediately below it as the
    caption; ``[[TABLE:n]]`` renders the markdown table that follows as a real
    Word table with the ``**Table n.**`` paragraph immediately below it.  There
    is no "Figure legends" section.  ``keep_with_next`` on the image paragraph
    and on every table row stops a page break from separating a display item
    from its caption.
4.  Bold: run-level bold is emitted **only** in a caption -- its ``**Figure n.**``
    / ``**Table n.**`` lead and its ``**a**`` panel letters.  Headings are bold
    through the Word heading styles, not through runs.  Every other ``**...**``
    span in the markdown renders as plain text; ``verify()`` walks the runs and
    asserts it.
5.  Emit two files: a clean copy, and a marked copy in which every run whose
    wording is not present verbatim in the corresponding v4 paragraph is
    highlighted yellow.  The correspondence is found by a word-level diff
    against the whole v4 paragraph list; because the revision is extensive, most
    of the document highlights, which is honest.
6.  Verify: reopen both, count paragraphs / figures / tables, audit bold, diff
    the clean text against the markdown, confirm each caption sits exactly one
    paragraph below its display item, resolve every cross-reference, render both
    to PDF and sweep for anything painted past the right margin.

Run:  /path/to/conda/bin/python build_manuscript_docx.py
      /path/to/conda/bin/python build_manuscript_docx.py --clean-only
      /path/to/conda/bin/python build_manuscript_docx.py --marked-only
"""

from __future__ import annotations

import array
import copy
import difflib
import os
import re
import shutil
import struct
import subprocess
import sys
import tempfile

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Emu, Inches, Pt
from docx.text.paragraph import Paragraph

# --------------------------------------------------------------------------
# paths
# --------------------------------------------------------------------------
HERE = os.path.dirname(os.path.abspath(__file__))
RESUB = os.path.dirname(HERE)                      # .../resubmission2
NC = os.path.dirname(RESUB)                        # .../nature_communications
V4_DOCX = os.path.join(NC, "1st_submission", "manuscript_v4.docx")
MD = os.path.join(RESUB, "manuscript_revised.md")
SUPP_MD = os.path.join(RESUB, "supplementary_revised.md")
FIGDIR = os.path.join(RESUB, "figures")
SUPPFIGDIR = os.path.join(RESUB, "supp_figures")
OUT_CLEAN = os.path.join(RESUB, "manuscript_revised.docx")
OUT_MARKED = os.path.join(RESUB, "manuscript_revised_marked.docx")

N_FIGURES = 5
N_TABLES = 1

# --------------------------------------------------------------------------
# page geometry (A4, 1 in margins) -- read back from the file, not assumed
# --------------------------------------------------------------------------
USABLE_TWIPS = None           # filled in from sectPr; 9026 for this document
TABLE_TOTAL_TWIPS = 8900      # < usable width, close to the v4 table's 9024
IMG_MAX_W_IN = 6.25
IMG_MAX_H_IN = 8.60

MARK_NOTE = (
    "Marked copy: yellow marks text new or rewritten since the original "
    "submission; unhighlighted text is unchanged and deletions are not shown."
)

# --------------------------------------------------------------------------
# markdown markers
# --------------------------------------------------------------------------
FIG_MARK_RE = re.compile(r"^\[\[FIGURE:(\d+)\]\]$")
TAB_MARK_RE = re.compile(r"^\[\[TABLE:(\d+)\]\]$")
CAPTION_RE = re.compile(r"^\*\*(Figure|Table)\s+(\d+)\.")


# ==========================================================================
# inline markdown -> runs
# ==========================================================================

ENDASH_RE = re.compile(r"--")
# pandoc rule: a backslash escapes any ASCII punctuation character
ESCAPABLE = set("!\"#$%&'()*+,-./:;<=>?@[\\]^_`{|}~")


def parse_inline(text: str) -> list:
    """Return ``[(text, bold, italic, code), ...]``.

    ``**`` toggles bold, ``*`` toggles italic, backticks delimit a code span,
    ``\\x`` is a literal ``x``, and ``--`` becomes an en dash everywhere except
    inside a code span (where it is a command-line flag).  An unbalanced marker
    degrades to a literal character rather than raising.
    """
    runs, buf = [], []
    bold = italic = False
    i, n = 0, len(text)

    def flush(code=False):
        if buf:
            t = "".join(buf)
            if not code:
                t = ENDASH_RE.sub("\u2013", t)
            runs.append((t, bold, italic, code))
            buf.clear()

    while i < n:
        ch = text[i]
        if ch == "\\" and i + 1 < n and text[i + 1] in ESCAPABLE:
            buf.append(text[i + 1])
            i += 2
            continue
        if ch == "`":
            j = text.find("`", i + 1)
            if j == -1:
                buf.append(ch)
                i += 1
                continue
            flush()
            code_txt = text[i + 1:j]
            if code_txt:
                runs.append((code_txt, bold, italic, True))
            i = j + 1
            continue
        if text.startswith("**", i):
            flush()
            bold = not bold
            i += 2
            continue
        if ch == "*":
            flush()
            italic = not italic
            i += 1
            continue
        buf.append(ch)
        i += 1
    flush()
    return [r for r in runs if r[0]]


def plain_text(text: str) -> str:
    """The rendered text of a markdown fragment: markers gone, escapes and
    en dashes resolved.  This is what ends up in the Word file."""
    return "".join(r[0] for r in parse_inline(text))


# ==========================================================================
# markdown block parsing
# ==========================================================================

def split_row(line: str) -> list:
    """Split a markdown table row on unescaped pipes."""
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    out, buf, i = [], [], 0
    while i < len(line):
        ch = line[i]
        if ch == "\\" and i + 1 < len(line) and line[i + 1] == "|":
            buf.append("|")
            i += 2
            continue
        if ch == "|":
            out.append("".join(buf).strip())
            buf = []
            i += 1
            continue
        buf.append(ch)
        i += 1
    out.append("".join(buf).strip())
    return out


SEP_RE = re.compile(r"^\|[\s:\-|]+\|?\s*$")


def parse_markdown(path: str) -> list:
    """Return a flat list of blocks: ``(kind, payload)``.

    ``kind`` is ``h`` (level, text), ``p`` (text), ``table`` (rows),
    ``figure`` (n) or ``tablemark`` (n).
    """
    lines = open(path, encoding="utf-8").read().split("\n")
    blocks, i, n = [], 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()

        if not s:
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append(("h", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue

        m = FIG_MARK_RE.match(s)
        if m:
            blocks.append(("figure", int(m.group(1))))
            i += 1
            continue

        m = TAB_MARK_RE.match(s)
        if m:
            blocks.append(("tablemark", int(m.group(1))))
            i += 1
            continue

        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                raw = lines[i].strip()
                if not SEP_RE.match(raw):
                    rows.append(split_row(raw))
                i += 1
            if rows:
                ncols = max(len(r) for r in rows)
                rows = [r + [""] * (ncols - len(r)) for r in rows]
                blocks.append(("table", rows))
            continue

        buf = [s]
        i += 1
        while i < n:
            nxt = lines[i]
            if not nxt.strip():
                break
            if nxt.strip().startswith(("|", "#")):
                break
            if FIG_MARK_RE.match(nxt.strip()) or TAB_MARK_RE.match(nxt.strip()):
                break
            buf.append(nxt.strip())
            i += 1
        blocks.append(("p", " ".join(buf)))
    return blocks


# ==========================================================================
# word-level diff, for the marked copy
# ==========================================================================

WORD_RE = re.compile(r"\S+")
MIN_EQUAL_RUN = 4        # shorter islands of unchanged words are absorbed
MIN_SIMILARITY = 0.35    # below this the paragraph counts as fully rewritten


def changed_ranges(old: str, new: str):
    """Character ranges of ``new`` that are not restored verbatim from ``old``.

    ``None`` means "everything is new"; ``[]`` means "nothing is new".
    """
    old = " ".join(old.split())
    new_norm = new
    a = [m.group(0) for m in WORD_RE.finditer(old)]
    spans = [(m.start(), m.end()) for m in WORD_RE.finditer(new_norm)]
    b = [new_norm[s:e] for s, e in spans]
    if not b:
        return []
    if not a:
        return None

    sm = difflib.SequenceMatcher(a=a, b=b, autojunk=False)
    if sm.ratio() < MIN_SIMILARITY:
        return None

    changed = [True] * len(b)
    for tag, _i1, _i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            for j in range(j1, j2):
                changed[j] = False

    # absorb tiny islands of unchanged words inside rewritten stretches
    k = 0
    while k < len(changed):
        if changed[k]:
            k += 1
            continue
        j = k
        while j < len(changed) and not changed[j]:
            j += 1
        if (j - k) < MIN_EQUAL_RUN and k > 0 and j < len(changed):
            for t in range(k, j):
                changed[t] = True
        k = j

    if all(changed):
        return None
    if not any(changed):
        return []

    ranges = []
    for j, flag in enumerate(changed):
        if not flag:
            continue
        s, e = spans[j]
        if ranges and new_norm[ranges[-1][1]:s].strip() == "":
            ranges[-1] = (ranges[-1][0], e)
        else:
            ranges.append((s, e))
    return ranges


class V4Wording:
    """The wording of the submitted manuscript, for the marked copy.

    ``match(text)`` returns the v4 paragraph (or table cell) whose wording is
    closest to ``text``, or ``None`` when nothing in v4 is close enough -- which
    is how a genuinely new paragraph is recognised.
    """

    def __init__(self, texts):
        self.texts = [" ".join(t.split()) for t in texts if t and t.strip()]
        self.words = [t.split() for t in self.texts]
        self._cache = {}

    def match(self, text):
        key = text
        if key in self._cache:
            return self._cache[key]
        b = text.split()
        best, best_r = None, MIN_SIMILARITY
        if b:
            for t, a in zip(self.texts, self.words):
                sm = difflib.SequenceMatcher(a=a, b=b, autojunk=False)
                if sm.real_quick_ratio() <= best_r:
                    continue
                if sm.quick_ratio() <= best_r:
                    continue
                r = sm.ratio()
                if r > best_r:
                    best_r, best = r, t
        self._cache[key] = best
        return best

    def ranges(self, text):
        """``hl_ranges`` for ``fill_runs``: ``None`` = mark the whole thing."""
        old = self.match(text)
        if old is None:
            return None
        return changed_ranges(old, text)


# ==========================================================================
# docx helpers
# ==========================================================================

CODE_FONT = "Consolas"
ZWSP = "\u200b"
BREAK_AFTER = set("/_-,.:;{}()=|\\&+")
LONG_TOKEN = 14


# --- tokens soft_break() must leave alone ---------------------------------
# A zero-width space is invisible on the page but it is a real character in
# the text stream: it survives copy-paste, and journal production converts
# Word to XML, where stray format characters are a known source of
# corruption.  Two classes of token are therefore exempt.
_OPENERS = "([{<\u2018\u201c\"'"
_CLOSERS = ")]}>\u2019\u201d\"'.,;:!?"
_URL_PREFIXES = ("http://", "https://", "www.", "doi:")
_DOI_RE = re.compile(r"10\.\d{4,9}/")   # a real DOI prefix, not "10.59/54.42"
_HYPHEN_WORD_RE = re.compile(r"[A-Za-z0-9]+(?:-[A-Za-z0-9]+)+")


def _is_link(tok: str) -> bool:
    """A web address or DOI, however prose wrapped it in brackets or quotes.

    The Data availability and Code availability statements exist so that a
    reader can follow the link; a hidden character inside one stops it
    resolving once it is copied out of the PDF.
    """
    core = tok.lstrip(_OPENERS).lower()
    return core.startswith(_URL_PREFIXES) or _DOI_RE.match(core) is not None


def _is_hyphenated_word(tok: str) -> bool:
    """``cluster-bootstrap``, ``MAGICC-minus-CheckM2``: alphanumeric parts
    joined by hyphens.  Word and LibreOffice break after a hyphen natively, so
    the helper would only add a hidden duplicate of a break opportunity the
    text already carries."""
    return _HYPHEN_WORD_RE.fullmatch(tok.strip(_OPENERS + _CLOSERS)) is not None


def soft_break(text: str) -> str:
    """Insert zero-width spaces inside over-long tokens.

    Result-file paths and accession lists are single unbreakable words; Word
    and LibreOffice both push them past the right margin rather than breaking
    them.  A zero-width space is an invisible break opportunity that both
    honour, and it does not change the visible text.
    """
    out = []
    for tok in re.split(r"(\s+)", text):
        if (len(tok) > LONG_TOKEN and not tok.isspace()
                and not _is_link(tok) and not _is_hyphenated_word(tok)):
            buf, run_len = [], 0
            for ch in tok:
                buf.append(ch)
                run_len += 1
                if ch in BREAK_AFTER and run_len >= 6:
                    buf.append(ZWSP)
                    run_len = 0
            tok = "".join(buf)
        out.append(tok)
    return "".join(out)


def set_run_font(run, name):
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)
    rf.attrib.pop(qn("w:hint"), None)


def set_italic_off(run):
    """Bibliography inherits an italic from its style in some Word builds; v4
    turned it off explicitly on every reference run, so we do the same."""
    rpr = run._r.get_or_add_rPr()
    for tag in ("w:i", "w:iCs"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
        el = OxmlElement(tag)
        el.set(qn("w:val"), "0")
        rpr.append(el)


def _split_by_ranges(text, offset, ranges):
    """Yield ``(fragment, highlighted)`` for ``text`` starting at ``offset``."""
    if ranges is None:
        yield text, True
        return
    if not ranges:
        yield text, False
        return
    pos = 0
    n = len(text)
    for rs, re_ in ranges:
        if re_ <= offset or rs >= offset + n:
            continue
        s = max(rs - offset, pos)
        e = min(re_ - offset, n)
        if s > pos:
            yield text[pos:s], False
        if e > s:
            yield text[s:e], True
        pos = max(pos, e)
    if pos < n:
        yield text[pos:], False


AUTOLINK_RE = re.compile(r"<((?:https?|ftp|mailto):[^>\s]+)>")


def fill_runs(par, text, size_pt=None, allow_bold=False, italic_off=False,
              hl=False, hl_ranges=()):
    """Emit formatted runs parsed from markdown into an empty paragraph.

    ``allow_bold`` is the whole bold policy: with it False -- everywhere except
    a caption -- a ``**...**`` span in the markdown renders as plain text.
    ``hl`` marks the whole paragraph; ``hl_ranges`` (character ranges over the
    rendered plain text) marks only part of it.  ``None`` ranges mean "all", the
    default empty tuple means "none", so the clean copy is never highlighted.
    """
    text = AUTOLINK_RE.sub(r"\1", text)
    for r in list(par._p.findall(qn("w:r"))):
        par._p.remove(r)
    for h in list(par._p.findall(qn("w:hyperlink"))):
        par._p.remove(h)

    offset = 0
    for txt, bold, italic, code in parse_inline(text):
        if hl:
            pieces = [(txt, True)]
        else:
            pieces = list(_split_by_ranges(txt, offset, hl_ranges))
        for frag, marked in pieces:
            run = par.add_run(soft_break(frag))
            run.bold = True if (bold and allow_bold) else None
            run.italic = True if italic else None
            if italic_off and not italic:
                set_italic_off(run)
            if code:
                set_run_font(run, CODE_FONT)
            if size_pt is not None:
                run.font.size = Pt(size_pt)
            if marked:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        offset += len(txt)
    return par


def new_paragraph(doc, style=None):
    p = OxmlElement("w:p")
    par = Paragraph(p, doc._body)
    if style is not None:
        par.style = doc.styles[style]
    return par


def set_pPr_bits(par, jc=None, before=None, after=None, line=None,
                 keep_next=None):
    ppr = par._p.get_or_add_pPr()
    if before is not None or after is not None or line is not None:
        sp = ppr.find(qn("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing")
            ppr.insert(0, sp)
        if before is not None:
            sp.set(qn("w:before"), str(before))
        if after is not None:
            sp.set(qn("w:after"), str(after))
        if line is not None:
            sp.set(qn("w:line"), str(line))
            sp.set(qn("w:lineRule"), "auto")
    if jc is not None:
        for j in list(ppr.findall(qn("w:jc"))):
            ppr.remove(j)
        el = OxmlElement("w:jc")
        el.set(qn("w:val"), jc)
        ppr.append(el)
    if keep_next:
        for k in list(ppr.findall(qn("w:keepNext"))):
            ppr.remove(k)
        ppr.insert(0, OxmlElement("w:keepNext"))
    return par


# ==========================================================================
# table metrics and layout
# ==========================================================================

# Arial's advance widths are Helvetica's; using the real table instead of an
# average keeps caps-heavy headers such as "CheckM2" from wrapping in a column
# that a mean-width estimate said was wide enough.
_HELV = {
    " ": 278, "!": 278, '"': 355, "#": 556, "$": 556, "%": 889, "&": 667,
    "'": 191, "(": 333, ")": 333, "*": 389, "+": 584, ",": 278, "-": 333,
    ".": 278, "/": 278, "0": 556, "1": 556, "2": 556, "3": 556, "4": 556,
    "5": 556, "6": 556, "7": 556, "8": 556, "9": 556, ":": 278, ";": 278,
    "<": 584, "=": 584, ">": 584, "?": 556, "@": 1015, "A": 667, "B": 667,
    "C": 722, "D": 722, "E": 667, "F": 611, "G": 778, "H": 722, "I": 278,
    "J": 500, "K": 667, "L": 556, "M": 833, "N": 722, "O": 778, "P": 667,
    "Q": 778, "R": 722, "S": 667, "T": 611, "U": 722, "V": 667, "W": 944,
    "X": 667, "Y": 667, "Z": 611, "[": 278, "\\": 278, "]": 278, "^": 469,
    "_": 556, "`": 333, "a": 556, "b": 556, "c": 500, "d": 556, "e": 556,
    "f": 278, "g": 556, "h": 556, "i": 222, "j": 222, "k": 500, "l": 222,
    "m": 833, "n": 556, "o": 556, "p": 556, "q": 556, "r": 333, "s": 500,
    "t": 278, "u": 556, "v": 500, "w": 722, "x": 500, "y": 500, "z": 500,
    "{": 334, "|": 260, "}": 334, "~": 584,
}
_DEFAULT_ADV = 600      # non-ASCII: en/em dash, +-, >=, superscripts, Greek
_BOLD_FACTOR = 1.07
_MONO_ADV = 550         # Consolas is 0.55 em for every glyph


def _adv(ch, bold, code):
    if code:
        return _MONO_ADV / 1000.0
    a = _HELV.get(ch, _DEFAULT_ADV) / 1000.0
    return a * _BOLD_FACTOR if bold else a


def md_metrics(cell_md, header=False):
    """(total width, longest unbreakable token) of a markdown cell, in em."""
    total, longest, cur = 0.0, 0.0, 0.0
    for txt, bold, _it, code in parse_inline(cell_md):
        b = bold or header
        for ch in txt:
            w = _adv(ch, b, code)
            total += w
            if ch.isspace():
                cur = 0.0
            else:
                cur += w
                longest = max(longest, cur)
    return total, longest


def text_em(text, bold=False):
    return sum(_adv(ch, bold, False) for ch in text)


CELL_PAD = 90            # twips: cell margins plus the two cell borders
MIN_COL_TWIPS = 480
TARGET_LINE_EM = 17.0    # a comfortable measure for a table cell, in em


def column_metrics(rows, size_pt):
    """(need, want) per column, in twips, at ``size_pt``."""
    ncols = len(rows[0])
    twips_per_em = size_pt * 20.0
    need, want = [], []
    for c in range(ncols):
        totals, longests = [], []
        for ri, r in enumerate(rows):
            t, l = md_metrics(r[c], header=(ri == 0))
            totals.append(t)
            longests.append(l)
        need.append(max(max(longests) * twips_per_em + CELL_PAD, MIN_COL_TWIPS))
        want.append(min(max(totals), TARGET_LINE_EM) * twips_per_em + CELL_PAD)
    return need, want


def allocate(total, need, want):
    """Split ``total`` proportionally to ``want`` but never below ``need``."""
    n = len(need)
    if sum(need) >= total:
        scale = total / sum(need)
        w = [x * scale for x in need]
    else:
        w = [None] * n
        free = set(range(n))
        remaining = float(total)
        while free:
            tot_w = sum(want[i] for i in free) or 1.0
            pinned = [i for i in free if remaining * want[i] / tot_w < need[i]]
            if not pinned:
                for i in free:
                    w[i] = remaining * want[i] / tot_w
                break
            for i in pinned:
                w[i] = need[i]
                remaining -= need[i]
                free.discard(i)
    out = [int(round(x)) for x in w]
    k = max(range(n), key=lambda i: out[i])
    out[k] += total - sum(out)
    return out


def table_layout(rows, total=TABLE_TOTAL_TWIPS):
    """Choose the body font size and the column widths together."""
    ncols = len(rows[0])
    max_row_em = max(sum(md_metrics(c, header=(ri == 0))[0] for c in r)
                     for ri, r in enumerate(rows))
    longest_cell = max(len(plain_text(c)) for r in rows for c in r)
    if ncols >= 6:
        start = 8.0
    elif ncols == 5 or longest_cell > 60:
        start = 9.0
    else:
        start = 10.0
    if max_row_em > 145 and start > 8.0:
        start -= 1.0
    if max_row_em > 235:
        start = 8.0

    for size_pt in [x for x in (10.0, 9.0, 8.0, 7.0, 6.5) if x <= start]:
        need, want = column_metrics(rows, size_pt)
        if sum(need) <= total:
            return size_pt, allocate(total, need, want)
    size_pt = 6.5
    need, want = column_metrics(rows, size_pt)
    scale = total / sum(need)
    need = [x * scale for x in need]
    return size_pt, allocate(total, need, want)


def build_table(doc, tbl_el, rows, wording=None, keep_next=True):
    """Rebuild ``tbl_el`` (a clone of the v4 ``w:tbl``) from ``rows``.

    Row and cell formatting come from the element's own first (header) and
    second (body) rows, so the table keeps the v4 borders, header shading and
    cell margins.  ``keep_next`` puts ``w:keepNext`` on every cell paragraph, so
    the table cannot be split from the legend that follows it.
    """
    trs = tbl_el.findall(qn("w:tr"))
    hdr_tpl = copy.deepcopy(trs[0])
    body_tpl = copy.deepcopy(trs[1]) if len(trs) > 1 else copy.deepcopy(trs[0])
    for tr in trs:
        tbl_el.remove(tr)

    def strip_to_one_cell(tr_tpl):
        tcs = tr_tpl.findall(qn("w:tc"))
        keep = copy.deepcopy(tcs[0])
        for tc in tr_tpl.findall(qn("w:tc")):
            tr_tpl.remove(tc)
        return tr_tpl, keep

    hdr_row_tpl, hdr_cell_tpl = strip_to_one_cell(hdr_tpl)
    body_row_tpl, body_cell_tpl = strip_to_one_cell(body_tpl)

    ncols = len(rows[0])
    size_pt, widths = table_layout(rows)

    aligns = []
    for c in range(ncols):
        longest = max(len(plain_text(r[c])) for r in rows)
        aligns.append("center" if longest <= 20 else "left")

    tblpr = tbl_el.find(qn("w:tblPr"))
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.insert(0, tblw)
    tblw.set(qn("w:w"), str(TABLE_TOTAL_TWIPS))
    tblw.set(qn("w:type"), "dxa")
    for old in tblpr.findall(qn("w:tblLayout")):
        tblpr.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    look = tblpr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(layout)
    else:
        tblpr.append(layout)

    grid = tbl_el.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tblpr.addnext(grid)
    for gc in list(grid.findall(qn("w:gridCol"))):
        grid.remove(gc)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)

    for ri, row in enumerate(rows):
        is_hdr = ri == 0
        tr = copy.deepcopy(hdr_row_tpl if is_hdr else body_row_tpl)
        cell_tpl = hdr_cell_tpl if is_hdr else body_cell_tpl
        for ci, cell_md in enumerate(row):
            tc = copy.deepcopy(cell_tpl)
            tcpr = tc.find(qn("w:tcPr"))
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.insert(0, tcw)
            tcw.set(qn("w:w"), str(widths[ci]))
            tcw.set(qn("w:type"), "dxa")
            ps = tc.findall(qn("w:p"))
            for extra in ps[1:]:
                tc.remove(extra)
            par = Paragraph(ps[0], doc._body)
            hl_ranges = []
            if wording is not None:
                hl_ranges = wording.ranges(plain_text(cell_md))
            fill_runs(par, cell_md, size_pt=size_pt, hl_ranges=hl_ranges)
            ppr = par._p.get_or_add_pPr()
            for j in list(ppr.findall(qn("w:jc"))):
                ppr.remove(j)
            jc = OxmlElement("w:jc")
            jc.set(qn("w:val"), "center" if is_hdr else aligns[ci])
            ppr.append(jc)
            if keep_next:
                for k in list(ppr.findall(qn("w:keepNext"))):
                    ppr.remove(k)
                ppr.insert(0, OxmlElement("w:keepNext"))
            tr.append(tc)
        tbl_el.append(tr)
    return tbl_el


# ==========================================================================
# images
# ==========================================================================

_TRIM_DIR = None
_TRIM_REPORT = []


def png_size(path):
    with open(path, "rb") as fh:
        head = fh.read(32)
    assert head[:8] == b"\x89PNG\r\n\x1a\n", path
    w, h = struct.unpack(">II", head[16:24])
    return w, h


def trimmed_png(src, pad_frac=0.004, level=245):
    """A copy of ``src`` with its uniform near-white border cropped away.

    ``figures/Figure_2.png`` is a schematic drawn on a canvas roughly 15 % of
    whose height is blank at the bottom; embedded untrimmed it reads as a gap
    between the figure and its caption.  The source file is never modified: the
    crop is written to a temporary directory that is removed when the build
    finishes.  A small pad is kept so no glyph touches the crop edge.
    """
    global _TRIM_DIR
    from PIL import Image

    if _TRIM_DIR is None:
        _TRIM_DIR = tempfile.mkdtemp(prefix="msfig_")
    dst = os.path.join(_TRIM_DIR, os.path.basename(src))
    if os.path.exists(dst):
        return dst

    im = Image.open(src)
    w, h = im.size
    data = array.array("B", im.convert("L").tobytes())
    # a row (column) is blank when every one of its pixels is near-white
    rows_ink = [y for y in range(h) if min(data[y * w:(y + 1) * w]) < level]
    cols_ink = [x for x in range(w) if min(data[x::w]) < level]
    if not rows_ink or not cols_ink:
        shutil.copyfile(src, dst)
        _TRIM_REPORT.append((os.path.basename(src), (w, h), (w, h)))
        return dst

    pad = int(round(min(w, h) * pad_frac))
    x0 = max(cols_ink[0] - pad, 0)
    x1 = min(cols_ink[-1] + 1 + pad, w)
    y0 = max(rows_ink[0] - pad, 0)
    y1 = min(rows_ink[-1] + 1 + pad, h)
    out = im.crop((x0, y0, x1, y1))
    out.save(dst, dpi=im.info.get("dpi", (400, 400)))
    _TRIM_REPORT.append((os.path.basename(src), (w, h), out.size))
    return dst


def fit(path):
    w, h = png_size(path)
    width_in = IMG_MAX_W_IN
    height_in = width_in * h / w
    if height_in > IMG_MAX_H_IN:
        height_in = IMG_MAX_H_IN
        width_in = height_in * w / h
    return Inches(width_in), Inches(height_in)


def figure_png(n):
    p = os.path.join(FIGDIR, f"Figure_{n}.png")
    assert os.path.exists(p), p
    return trimmed_png(p)


_DOCPR_ID = [900000000]


def picture_paragraph(doc, template_p, png_path):
    """A new image paragraph built from a *clone of a v4 image paragraph*.

    ``run.add_picture`` emits a bare ``wp:inline`` with no ``distL``/``distR``
    and no ``effectExtent``.  LibreOffice lays such an inline shape out ~9 pt to
    the right of where it puts the v4 ones, which pushes wide figures past the
    right margin.  Cloning the v4 drawing keeps every attribute Word wrote; only
    the relationship id and the extent change.
    """
    rid, _part = doc.part.get_or_add_image(png_path)
    p_el = copy.deepcopy(template_p)
    par = Paragraph(p_el, doc._body)

    for br in p_el.findall(".//" + qn("w:lastRenderedPageBreak")):
        br.getparent().remove(br)
    for t in p_el.findall(".//" + qn("w:t")):
        t.getparent().remove(t)

    blips = p_el.findall(".//" + qn("a:blip"))
    assert len(blips) == 1
    blips[0].set(qn("r:embed"), rid)
    for ext in blips[0].findall(".//" + qn("a:ext")):
        ext.getparent().remove(ext)      # a14:useLocalDpi hint, now stale

    cx, cy = fit(png_path)
    for ext in p_el.findall(".//" + qn("wp:extent")):
        ext.set("cx", str(int(cx)))
        ext.set("cy", str(int(cy)))
    for ext in p_el.findall(".//" + qn("a:ext")):
        ext.set("cx", str(int(cx)))
        ext.set("cy", str(int(cy)))
    for eff in p_el.findall(".//" + qn("wp:effectExtent")):
        for a in ("l", "t", "r", "b"):
            eff.set(a, "0")
    name = os.path.basename(png_path)
    _DOCPR_ID[0] += 1
    for docpr in p_el.findall(".//" + qn("wp:docPr")):
        docpr.set("id", str(_DOCPR_ID[0]))
        docpr.set("name", name)
        docpr.set("descr", name)
        if "title" in docpr.attrib:
            docpr.set("title", name)
    return par


# ==========================================================================
# build
# ==========================================================================

BODY_AFTER = 120
CAPTION_PT = 10.0
SMALL_PT = 9.0


def build(out_path, marked):
    """Copy the pristine v4 file, empty its body, render the markdown into it."""
    global USABLE_TWIPS

    shutil.copyfile(V4_DOCX, out_path)
    doc = docx.Document(out_path)

    sect = doc.sections[0]
    USABLE_TWIPS = int(
        (sect.page_width - sect.left_margin - sect.right_margin) / 635)
    assert TABLE_TOTAL_TWIPS <= USABLE_TWIPS, (TABLE_TOTAL_TWIPS, USABLE_TWIPS)

    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    assert sectPr is not None

    # ---- furniture captured from v4 before the body is emptied -----------
    v4_paras = list(doc.paragraphs)
    v4_texts = [p.text for p in v4_paras]
    v4_cells = [c.text for r in doc.tables[0].rows for c in r.cells]
    before = dict(paragraphs=len(v4_paras), tables=len(doc.tables),
                  images=len(doc.inline_shapes))
    assert before == dict(paragraphs=110, tables=1, images=4), before

    pic_template = copy.deepcopy(
        next(p._p for p in v4_paras if p._p.findall(".//" + qn("a:blip"))))
    table_template = copy.deepcopy(doc.tables[0]._tbl)

    wording = V4Wording(v4_texts) if marked else None
    cell_wording = V4Wording(v4_cells) if marked else None

    # ---- empty the body, keep sectPr -------------------------------------
    for child in list(body.iterchildren()):
        if child is sectPr:
            continue
        body.remove(child)

    def emit(el):
        sectPr.addprevious(el)
        return el

    # ---- render ----------------------------------------------------------
    blocks = parse_markdown(MD)
    n_head = 0
    in_refs = False
    figures_placed, tables_placed, captions = [], [], []
    para_no = 0                       # index into doc.paragraphs as we build

    def para(style, text, size_pt=None, jc="both", before=None,
             after=BODY_AFTER, line=None, allow_bold=False, italic_off=False,
             keep_next=False):
        nonlocal para_no
        p = new_paragraph(doc, style)
        set_pPr_bits(p, jc=jc, before=before, after=after, line=line,
                     keep_next=keep_next)
        hl = []
        if wording is not None:
            hl = wording.ranges(plain_text(AUTOLINK_RE.sub(r"\1", text)))
        fill_runs(p, text, size_pt=size_pt, allow_bold=allow_bold,
                  italic_off=italic_off, hl_ranges=hl)
        emit(p._p)
        para_no += 1
        return p

    i = 0
    while i < len(blocks):
        kind, payload = blocks[i]

        # -------- headings --------------------------------------------------
        if kind == "h":
            level, text = payload
            style = "Heading 1" if level == 1 else "Heading 2"
            p = new_paragraph(doc, style)
            hl = []
            if wording is not None:
                hl = wording.ranges(plain_text(text))
            fill_runs(p, text, hl_ranges=hl)
            emit(p._p)
            para_no += 1
            n_head += 1
            in_refs = (level == 1 and plain_text(text).strip().lower()
                       == "references")
            i += 1
            continue

        # -------- figure ----------------------------------------------------
        if kind == "figure":
            n = payload
            png = figure_png(n)
            p = picture_paragraph(doc, pic_template, png)
            set_pPr_bits(p, jc="center", before=240, after=120, keep_next=True)
            emit(p._p)
            fig_idx = para_no
            para_no += 1
            figures_placed.append((n, fig_idx))

            kind2, cap_md = blocks[i + 1]
            m = CAPTION_RE.match(cap_md) if kind2 == "p" else None
            assert m and m.group(1) == "Figure" and int(m.group(2)) == n, (
                f"[[FIGURE:{n}]] is not followed by its '**Figure {n}.**' "
                f"caption: {str(cap_md)[:80]!r}")
            para("Normal", cap_md, size_pt=CAPTION_PT, jc="both",
                 before=60, after=240, allow_bold=True)
            captions.append(("Figure", n, fig_idx, fig_idx + 1))
            i += 2
            continue

        # -------- table -----------------------------------------------------
        if kind == "tablemark":
            n = payload
            kind2, rows = blocks[i + 1]
            assert kind2 == "table", f"[[TABLE:{n}]] is not followed by a table"
            tbl_el = copy.deepcopy(table_template)
            build_table(doc, tbl_el, rows, wording=cell_wording)
            emit(tbl_el)
            tab_idx = para_no          # the legend is the next paragraph
            tables_placed.append((n, len(rows), len(rows[0])))

            kind3, cap_md = blocks[i + 2]
            m = CAPTION_RE.match(cap_md) if kind3 == "p" else None
            assert m and m.group(1) == "Table" and int(m.group(2)) == n, (
                f"[[TABLE:{n}]] is not followed by its '**Table {n}.**' legend")
            para("Normal", cap_md, size_pt=CAPTION_PT, jc="both",
                 before=60, after=240, allow_bold=True)
            captions.append(("Table", n, None, tab_idx))
            i += 3
            continue

        if kind == "table":
            raise AssertionError(
                "a markdown table with no [[TABLE:n]] marker before it")

        # -------- paragraphs -------------------------------------------------
        text = payload
        assert not CAPTION_RE.match(text), (
            f"caption not attached to a display item: {text[:70]!r}")

        if n_head == 0:
            # front matter: title, authors, affiliations
            if para_no == 0:
                para("Title", text, jc=None, after=None)
            elif para_no == 1:
                para("Normal", text)
            else:
                para("Normal", text, size_pt=SMALL_PT)
        elif in_refs:
            para("Bibliography", text, size_pt=SMALL_PT, jc="both", line=240,
                 after=None, italic_off=True)
        else:
            para("Normal", text)
        i += 1

    # ---- the marked copy explains its own convention ----------------------
    if marked:
        note = new_paragraph(doc, "Normal")
        set_pPr_bits(note, jc="both", after=240)
        run = note.add_run(MARK_NOTE)
        run.italic = True
        run.font.size = Pt(9)
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        first = body.find(qn("w:p"))
        first.addprevious(note._p)

    # ---- drop the four v4 image parts, now unreferenced -------------------
    used = {b.get(qn("r:embed")) for b in body.findall(".//" + qn("a:blip"))}
    used |= {b.get(qn("r:link")) for b in body.findall(".//" + qn("a:blip"))}
    dropped = 0
    for rId, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/image") and rId not in used:
            try:
                doc.part.drop_rel(rId)
            except Exception:            # older Relationships mappings
                doc.part.rels._rels.pop(rId, None)
            dropped += 1

    doc.save(out_path)

    d2 = docx.Document(out_path)
    after = dict(paragraphs=len(d2.paragraphs), tables=len(d2.tables),
                 images=len(d2.inline_shapes))
    return dict(before=before, after=after, figures=figures_placed,
                tables=tables_placed, captions=captions, dropped_rels=dropped)


# ==========================================================================
# verification
# ==========================================================================

QUOTES = str.maketrans({"\u2019": "'", "\u2018": "'", "\u201c": '"',
                        "\u201d": '"', "\u2011": "-", "\u2010": "-"})


def norm(s):
    """Whitespace-collapsed text with the two quote shapes folded together."""
    return " ".join(s.replace(ZWSP, "").translate(QUOTES).split())


def target_paragraphs():
    """The paragraph texts ``manuscript_revised.md`` asks for, rendered the way
    the Word file renders them.  Tables and the ``[[FIGURE:n]]`` /
    ``[[TABLE:n]]`` position markers are dropped: they are objects, not
    paragraphs.  Pandoc ``<url>`` autolinks lose their angle brackets, as Word
    renders the bare URL."""
    out = []
    for kind, payload in parse_markdown(MD):
        if kind == "h":
            out.append(norm(AUTOLINK_RE.sub(r"\1", plain_text(payload[1]))))
        elif kind == "p":
            out.append(norm(AUTOLINK_RE.sub(r"\1", plain_text(payload))))
    return out


def docx_paragraphs(doc):
    out = []
    for p in doc.paragraphs:
        if p._p.findall(".//" + qn("a:blip")):
            continue
        t = norm(p.text)
        if t:
            out.append(t)
    return out


MAIN_FIG_RE = re.compile(r"\bFig(?:ure|s|\.)?\s*(?!S)(\d+)")
MAIN_TAB_RE = re.compile(r"\bTable\s*(?!S)(\d+)")
SUPP_FIG_RE = re.compile(r"\bFig(?:ures?|s?\.)\s+((?:S\d+[a-z]?"
                         r"(?:\s*[\u2013\u2014-]\s*S?\d+[a-z]?)?"
                         r"(?:\s*(?:,|and)\s*)?)+)")
SUPP_TAB_RE = re.compile(r"\bTables?\s+((?:S\d+[a-z]?"
                         r"(?:\s*[\u2013\u2014-]\s*S?\d+[a-z]?)?"
                         r"(?:\s*(?:,|and)\s*)?)+)")
SNUM_RE = re.compile(r"S?(\d+)")


def cross_references(paras):
    figs, tabs = set(), set()
    for t in paras:
        for m in MAIN_FIG_RE.finditer(t):
            figs.add(int(m.group(1)))
        for m in MAIN_TAB_RE.finditer(t):
            tabs.add(int(m.group(1)))
    return figs, tabs


def supp_references(paras):
    """Supplementary figure / table numbers cited anywhere in the text.

    ``Figs. S1-S5`` expands to S1..S5; ``Tables S6, S8`` gives S6 and S8.
    """
    figs, tabs = set(), set()
    for t in paras:
        for rx, sink in ((SUPP_FIG_RE, figs), (SUPP_TAB_RE, tabs)):
            for m in rx.finditer(t):
                frag = m.group(1)
                for part in re.split(r"\s*(?:,|and)\s*", frag):
                    part = part.strip()
                    if not part:
                        continue
                    rng = re.split(r"\s*[\u2013\u2014-]\s*", part)
                    nums = [int(SNUM_RE.match(x).group(1)) for x in rng
                            if SNUM_RE.match(x)]
                    if len(nums) == 2 and nums[1] >= nums[0]:
                        sink.update(range(nums[0], nums[1] + 1))
                    elif nums:
                        sink.add(nums[0])
    return figs, tabs


def supp_items_available():
    """Supplementary figures / tables that actually exist in the package."""
    figs = set()
    for name in os.listdir(SUPPFIGDIR):
        m = re.fullmatch(r"Figure_S(\d+)\.png", name)
        if m:
            figs.add(int(m.group(1)))
    tabs = set()
    if os.path.exists(SUPP_MD):
        txt = open(SUPP_MD, encoding="utf-8").read()
        for m in re.finditer(r"^#\s+Table S(\d+)\b", txt, re.M):
            tabs.add(int(m.group(1)))
    return figs, tabs


def bold_audit(doc):
    """Every run-level-bold run in the document, with where it sits."""
    hits = []
    for i, p in enumerate(doc.paragraphs):
        for r in p.runs:
            if r.bold:
                hits.append(("paragraph", i, p.style.name, norm(p.text)[:60],
                             r.text))
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            for ci, c in enumerate(row.cells):
                for p in c.paragraphs:
                    for r in p.runs:
                        if r.bold:
                            hits.append(("table", (ti, ri, ci), "-",
                                         norm(c.text)[:40], r.text))
    return hits


def pdf_check(path, label):
    tmp = tempfile.mkdtemp(prefix="msdocx_")
    r = subprocess.run(["/usr/bin/libreoffice", "--headless", "--convert-to",
                        "pdf", "--outdir", tmp, path],
                       capture_output=True, text=True, timeout=1800)
    pdf = os.path.join(tmp, os.path.basename(path).replace(".docx", ".pdf"))
    if not os.path.exists(pdf):
        print("  PDF CONVERSION FAILED:", r.stdout[-800:], r.stderr[-800:])
        raise SystemExit(1)
    info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True)
    m = re.search(r"^Pages:\s+(\d+)", info.stdout, re.M)
    pages = int(m.group(1)) if m else -1
    ps = re.search(r"^Page size:\s+(.*)$", info.stdout, re.M)
    size = ps.group(1) if ps else "?"
    print(f"  [{label}] PDF rendered OK: {os.path.getsize(pdf) / 1e6:.2f} MB, "
          f"{pages} pages, page size {size}")
    dims = re.match(r"([\d.]+)\s*x\s*([\d.]+)\s*pts", size)
    assert dims, f"{label}: cannot read the page size -- {size}"
    wpt, hpt = round(float(dims.group(1)), 1), round(float(dims.group(2)), 1)
    print(f"  [{label}] page size {wpt} x {hpt} pts "
          f"(A4 is 595.3 x 841.9): {'A4' if (wpt, hpt) == (595.3, 841.9) else 'NOT A4'}")
    assert (wpt, hpt) == (595.3, 841.9), f"{label}: not A4 -- {size}"

    bbox = subprocess.run(["pdftotext", "-bbox", pdf, "-"],
                          capture_output=True, text=True).stdout
    pw = 595.276                      # A4 width in points
    right_limit = pw - 72.0 + 1.5     # 1 in margin + 1.5 pt tolerance
    page_no, offenders = 0, []
    for line in bbox.split("\n"):
        if "<page " in line:
            page_no += 1
        m = re.search(r'<word xMin="([\d.]+)" yMin="([\d.]+)" '
                      r'xMax="([\d.]+)" yMax="([\d.]+)">(.*)</word>', line)
        if not m:
            continue
        if float(m.group(3)) > right_limit:
            offenders.append((page_no, round(float(m.group(3)), 1),
                              m.group(5)[:40]))
    print(f"  [{label}] words painted past the right margin "
          f"({right_limit - 1.5:.0f} pt): {len(offenders)}")
    for o in offenders[:8]:
        print(f"      page {o[0]}  xMax {o[1]}  {o[2]!r}")

    # ink sweep: catches images and table rules, which carry no words
    ink = ink_past_margin(pdf)
    print(f"  [{label}] pages with ink past the right margin: {len(ink)}"
          + (f"  {ink[:6]}" if ink else ""))

    os.remove(pdf)
    shutil.rmtree(tmp, ignore_errors=True)
    assert not offenders, f"{label}: content overflows the page width"
    assert not ink, f"{label}: ink painted past the right margin on {ink}"
    return pages, size


def ink_past_margin(pdf, dpi=150, tol_pt=3.0, level=128):
    """Pages carrying a dark pixel well to the right of the 1 in margin."""
    tmp = tempfile.mkdtemp(prefix="msppm_")
    try:
        subprocess.run(["pdftoppm", "-r", str(dpi), "-gray", pdf,
                        os.path.join(tmp, "p")],
                       capture_output=True, check=True, timeout=1800)
        bad = []
        for name in sorted(os.listdir(tmp)):
            if not name.endswith(".pgm"):
                continue
            with open(os.path.join(tmp, name), "rb") as fh:
                data = fh.read()
            fields, pos = [], 2
            while len(fields) < 3:
                while data[pos:pos + 1].isspace():
                    pos += 1
                if data[pos:pos + 1] == b"#":
                    while data[pos:pos + 1] not in (b"\n", b""):
                        pos += 1
                    continue
                start = pos
                while not data[pos:pos + 1].isspace():
                    pos += 1
                fields.append(int(data[start:pos]))
            pos += 1
            w, h, _mx = fields
            raster = data[pos:]
            x0 = int(round((595.276 - 72.0 + tol_pt) / 72.0 * dpi))
            x0 = min(max(x0, 0), w - 1)
            page = int(re.search(r"(\d+)\.pgm$", name).group(1))
            for y in range(h):
                row = raster[y * w:(y + 1) * w]
                if any(px < level for px in row[x0:]):
                    bad.append(page)
                    break
        return sorted(set(bad))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def verify(res_clean, res_marked):
    print("=" * 76)
    print("VERIFICATION")
    print("=" * 76)
    b = res_clean["before"]
    print(f"  v4 furniture source: paragraphs={b['paragraphs']:>4}  "
          f"tables={b['tables']:>2}  inline shapes={b['images']:>2}")
    for label, path, res in (("clean ", OUT_CLEAN, res_clean),
                             ("marked", OUT_MARKED, res_marked)):
        a = res["after"]
        print(f"  {label}             : paragraphs={a['paragraphs']:>4}  "
              f"tables={a['tables']:>2}  inline shapes={a['images']:>2}"
              f"   ({os.path.getsize(path) / 1e6:.2f} MB)")
        assert a["images"] == N_FIGURES, (label, a["images"])
        assert a["tables"] == N_TABLES, (label, a["tables"])
    print(f"  OK  {N_FIGURES} inline shapes and {N_TABLES} Word table in both")
    print("  figure whitespace trim (source files untouched):")
    for name, old, new in sorted(set(_TRIM_REPORT)):
        cut = 100.0 * (1 - (new[0] * new[1]) / (old[0] * old[1]))
        print(f"    {name}: {old[0]}x{old[1]} -> {new[0]}x{new[1]}  "
              f"({cut:.1f}% of canvas area removed)")
    print(f"  unreferenced v4 image parts dropped: "
          f"{res_clean['dropped_rels']}")

    doc = docx.Document(OUT_CLEAN)

    # ---- text diff against manuscript_revised.md -------------------------
    want = target_paragraphs()
    got = docx_paragraphs(doc)
    sm = difflib.SequenceMatcher(a=want, b=got, autojunk=False)
    diffs = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        for k in range(max(i2 - i1, j2 - j1)):
            w = want[i1 + k] if i1 + k < i2 else None
            g = got[j1 + k] if j1 + k < j2 else None
            diffs.append((tag, w, g))
    print(f"  target paragraphs  : {len(want)} in manuscript_revised.md; "
          f"{len(got)} non-empty text paragraphs in the docx")
    print(f"  differing paragraphs against manuscript_revised.md: {len(diffs)}")
    for tag, w, g in diffs[:12]:
        print(f"    [{tag}]")
        print(f"      md  : {(w or '')[:160]!r}")
        print(f"      docx: {(g or '')[:160]!r}")
    assert not diffs, f"{len(diffs)} paragraphs differ from the target markdown"

    # ---- legend placement -------------------------------------------------
    all_paras = doc.paragraphs
    img_idx = [i for i, p in enumerate(all_paras)
               if p._p.findall(".//" + qn("a:blip"))]
    print(f"  figure paragraphs at indices: {img_idx}")
    ok = True
    for kind, n, item_idx, cap_idx in res_clean["captions"]:
        cap = norm(all_paras[cap_idx].text)
        good = cap.startswith(f"{kind} {n}.")
        if kind == "Figure":
            good = good and (cap_idx == item_idx + 1) and item_idx in img_idx
            where = (f"image paragraph {item_idx}, caption paragraph "
                     f"{cap_idx} (delta {cap_idx - item_idx})")
        else:
            where = f"w:tbl element, legend paragraph {cap_idx}"
        print(f"    {kind} {n}: {where}  starts {cap[:34]!r}  "
              f"{'OK' if good else 'BAD'}")
        ok = ok and good
    assert ok, "a caption is not immediately below its display item"
    # the Word table's legend is the paragraph that directly follows the w:tbl
    body_children = list(doc.element.body.iterchildren())
    for k, child in enumerate(body_children):
        if child.tag == qn("w:tbl"):
            nxt = body_children[k + 1]
            assert nxt.tag == qn("w:p"), "no paragraph after the table"
            txt = norm(Paragraph(nxt, doc._body).text)
            assert txt.startswith("Table 1."), txt[:60]
            print(f"    the w:tbl is directly followed by {txt[:34]!r}  OK")
    # no separate legends section
    heads = [norm(p.text).lower() for p in all_paras
             if p.style.name.startswith("Heading")]
    bad_heads = [h for h in heads if "legend" in h]
    print(f"  'Figure legends' / 'Table legends' sections: "
          f"{bad_heads or 'none'}")
    assert not bad_heads

    # ---- keep-with-next ---------------------------------------------------
    def has_keep(p_el):
        ppr = p_el.find(qn("w:pPr"))
        return ppr is not None and ppr.find(qn("w:keepNext")) is not None
    kept_imgs = sum(1 for i in img_idx if has_keep(all_paras[i]._p))
    tbl_cells = [p for t in doc.tables for row in t.rows for c in row.cells
                 for p in c.paragraphs]
    kept_cells = sum(1 for p in tbl_cells if has_keep(p._p))
    print(f"  keep_with_next: {kept_imgs}/{len(img_idx)} figure paragraphs, "
          f"{kept_cells}/{len(tbl_cells)} table cell paragraphs")
    assert kept_imgs == len(img_idx) == N_FIGURES
    assert kept_cells == len(tbl_cells)

    # ---- tables against the markdown -------------------------------------
    md_tables = [p for k, p in parse_markdown(MD) if k == "table"]
    assert len(md_tables) == len(doc.tables) == N_TABLES
    bad = 0
    for ti, (mt, dt) in enumerate(zip(md_tables, doc.tables)):
        assert (len(mt), len(mt[0])) == (len(dt.rows), len(dt.columns)), (
            ti, len(mt), len(mt[0]), len(dt.rows), len(dt.columns))
        for ri, row in enumerate(mt):
            for ci, cell in enumerate(row):
                w = norm(plain_text(cell))
                g = norm(dt.cell(ri, ci).text)
                if w != g:
                    bad += 1
                    if bad <= 8:
                        print(f"    CELL DIFF t{ti} r{ri} c{ci}: "
                              f"md={w!r} docx={g!r}")
    print(f"  table cell differences against the markdown: {bad}")
    assert bad == 0

    # ---- no raw markdown survives ----------------------------------------
    scan = [norm(p.text) for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                scan.append(norm(c.text))
    leftovers = {}
    # a lone asterisk is legitimate: it is the footnote marker on the
    # corresponding-author line.  Only a *paired* marker is unconverted markup.
    for pat, name in ((r"\*\*", "**"), (r"\*\S[^*\n]{0,300}?\S\*", "*...*"),
                      (r"^\s*\|.*\|\s*$", "| table row"),
                      (r"\|\s*-{3,}\s*\|", "|--- rule"),
                      (r"^#{1,6}\s", "#"), (r"\[\[FIGURE", "[[FIGURE"),
                      (r"\[\[TABLE", "[[TABLE"),
                      (r"\bTODO\b", "TODO"),
                      (r"\bTBD\b", "TBD"), (r"PLACEHOLDER", "PLACEHOLDER")):
        hits = [s for s in scan if re.search(pat, s, re.M)]
        if hits:
            leftovers[name] = hits[:3]
    print("  raw-markdown / TODO leftovers: "
          + (str({k: len(v) for k, v in leftovers.items()}) if leftovers
             else "none"))
    assert not leftovers, leftovers

    # ---- cross-references -------------------------------------------------
    body_all = [norm(p.text) for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                body_all.append(norm(c.text))
    figs, tabs = cross_references(body_all)
    figs = {n for n in figs if n <= 30}
    tabs = {n for n in tabs if n <= 30}
    print(f"  main figures cited: {sorted(figs)}")
    print(f"  main tables  cited: {sorted(tabs)}")
    missing_f = [n for n in range(1, N_FIGURES + 1) if n not in figs]
    missing_t = [n for n in range(1, N_TABLES + 1) if n not in tabs]
    stray_f = sorted(n for n in figs if n > N_FIGURES)
    stray_t = sorted(n for n in tabs if n > N_TABLES)
    print(f"  main figures never cited: {missing_f or 'none'}   "
          f"citations to a figure that does not exist: {stray_f or 'none'}")
    print(f"  main tables never cited : {missing_t or 'none'}   "
          f"citations to a table that does not exist: {stray_t or 'none'}")
    assert not missing_f and not missing_t and not stray_f and not stray_t

    sfigs, stabs = supp_references(body_all)
    have_f, have_t = supp_items_available()
    print(f"  supplementary figures cited: {sorted(sfigs)}")
    print(f"  supplementary figures present in supp_figures/: {sorted(have_f)}")
    print(f"  supplementary tables cited : {sorted(stabs)}")
    print(f"  supplementary tables defined in supplementary_revised.md: "
          f"{sorted(have_t)}")
    unresolved_f = sorted(sfigs - have_f)
    unresolved_t = sorted(stabs - have_t)
    print(f"  UNRESOLVED supplementary figure references: "
          f"{unresolved_f or 'none'}")
    print(f"  UNRESOLVED supplementary table references : "
          f"{unresolved_t or 'none'}")

    # ---- author-fillable placeholders ------------------------------------
    md_text = open(MD, encoding="utf-8").read()
    doc_text = "\n".join(scan)
    print("  author-fillable placeholders preserved:")
    for ph in ("[DOI]", "[RELEASE]", "[DATE]", "[ANALYSIS REPO / DOI]",
               "AUTHOR TO VERIFY"):
        n_md = md_text.count(ph)
        n_doc = doc_text.count(ph)
        print(f"    {ph:<22} markdown {n_md}   docx {n_doc}")
        assert n_doc == n_md, ph

    # ---- bold audit -------------------------------------------------------
    hits = bold_audit(doc)
    cap_idx = {c[3] for c in res_clean["captions"]}
    offenders = [h for h in hits
                 if not (h[0] == "paragraph" and h[1] in cap_idx)]
    by_para = {}
    for h in hits:
        if h[0] == "paragraph":
            by_para.setdefault(h[1], []).append(h[4])
    print(f"  run-level bold runs in the clean copy: {len(hits)}")
    for idx in sorted(by_para):
        runs = by_para[idx]
        print(f"    paragraph {idx} ({norm(doc.paragraphs[idx].text)[:28]!r}): "
              f"{len(runs)} bold runs -> {[r[:24] for r in runs][:8]}")
    print(f"  bold runs outside a caption: {len(offenders)}")
    for o in offenders[:8]:
        print(f"    {o}")
    assert not offenders, "bold outside a figure/table caption"
    hstyles = sorted({p.style.name for p in doc.paragraphs
                      if p.style.name.startswith("Heading")})
    n_head = sum(1 for p in doc.paragraphs
                 if p.style.name.startswith("Heading"))
    print(f"  headings carry their bold from the Word styles ({hstyles}): "
          f"{n_head} heading paragraphs, 0 run-level bold among them")
    assert not any(h[0] == "paragraph"
                   and doc.paragraphs[h[1]].style.name.startswith("Heading")
                   for h in hits)
    ital = sum(1 for p in doc.paragraphs for r in p.runs if r.italic)
    print(f"  real italic runs: {ital}")
    assert ital > 0

    # ---- table and image geometry ----------------------------------------
    for i, t in enumerate(doc.tables):
        g = t._tbl.find(qn("w:tblGrid"))
        w = sum(int(c.get(qn("w:w"))) for c in g.findall(qn("w:gridCol")))
        sizes = sorted({round(r.font.size.pt, 1) for row in t.rows
                        for c in row.cells for p in c.paragraphs
                        for r in p.runs if r.font.size})
        print(f"  Table {i + 1}: {len(t.rows)}x{len(t.columns)}, width {w} "
              f"twips ({w / 1440:.3f} in) of {USABLE_TWIPS} usable "
              f"({USABLE_TWIPS / 1440:.3f} in), body font {sizes} pt")
        assert w <= USABLE_TWIPS, (i, w)

    cell_over = []
    for ti, t in enumerate(doc.tables):
        g = t._tbl.find(qn("w:tblGrid"))
        cw = [int(c.get(qn("w:w"))) for c in g.findall(qn("w:gridCol"))]
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                if ci >= len(cw):
                    continue
                size = max((r.font.size.pt for p in cell.paragraphs
                            for r in p.runs if r.font.size), default=10.0)
                bold = any(r.bold for p in cell.paragraphs for r in p.runs)
                worst = max((text_em(x, bold) for x in
                             re.split(r"[\s\u200b]+", cell.text) if x),
                            default=0.0)
                if worst * size * 20.0 > cw[ci] - CELL_PAD + 20:
                    cell_over.append((ti, ri, ci, round(worst, 1),
                                      cell.text[:40]))
    print(f"  table cells with an unbreakable run wider than their column: "
          f"{len(cell_over)}")
    assert not cell_over, cell_over[:5]

    body_over = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=11.0)
        bold = any(r.bold for r in p.runs)
        worst = max((text_em(t, bold) for t in
                     re.split(r"[\s\u200b]+", p.text) if t), default=0.0)
        if worst * size * 20.0 > USABLE_TWIPS:
            body_over.append((round(worst, 1), p.text[:70]))
    print(f"  body paragraphs with an unbreakable run wider than the text "
          f"column: {len(body_over)}")
    for w, t in body_over[:5]:
        print(f"    {w} em: {t!r}")
    assert not body_over

    usable_h = int(doc.sections[0].page_height - doc.sections[0].top_margin
                   - doc.sections[0].bottom_margin)
    print("  inline shapes:")
    for i, s in enumerate(doc.inline_shapes, start=1):
        pw, ph = png_size(figure_png(i))
        ar, ar_want = s.width / s.height, pw / ph
        print(f"    Figure {i}: {Emu(s.width).inches:.3f} x "
              f"{Emu(s.height).inches:.3f} in   aspect {ar:.4f} "
              f"(trimmed png {ar_want:.4f})")
        assert abs(ar - ar_want) / ar_want < 0.01, (i, ar, ar_want)
        assert s.width <= USABLE_TWIPS * 635, (i, s.width)
        assert s.height <= usable_h, (i, s.height)
    print(f"  OK  every image <= {USABLE_TWIPS / 1440:.3f} in wide, aspect "
          f"ratio preserved, height <= {Emu(usable_h).inches:.2f} in")

    order = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p") and child.findall(".//" + qn("a:blip")):
            order.append(len(order) + 1)
    print(f"  figures in document order: {len(order)} images, Fig 1..{len(order)}")
    assert order == list(range(1, N_FIGURES + 1))

    # ---- the clean copy carries no marking -------------------------------
    stray = sum(1 for p in doc.paragraphs for r in p.runs
                if r.font.highlight_color is not None)
    stray += sum(1 for t in doc.tables for row in t.rows for c in row.cells
                 for p in c.paragraphs for r in p.runs
                 if r.font.highlight_color is not None)
    print(f"  highlighted runs in the clean copy: {stray}")
    assert stray == 0, "the clean copy must carry no highlighting"

    # ---- the marked copy -------------------------------------------------
    dm = docx.Document(OUT_MARKED)
    hl = sum(1 for p in dm.paragraphs for r in p.runs
             if r.font.highlight_color == WD_COLOR_INDEX.YELLOW)
    hl_t = sum(1 for t in dm.tables for row in t.rows for c in row.cells
               for p in c.paragraphs for r in p.runs
               if r.font.highlight_color == WD_COLOR_INDEX.YELLOW)
    unhl = sum(1 for p in dm.paragraphs for r in p.runs
               if r.font.highlight_color != WD_COLOR_INDEX.YELLOW)
    note_ok = MARK_NOTE.split(":")[0] in norm(dm.paragraphs[0].text)
    print(f"  marked copy: {hl} highlighted body runs, {hl_t} highlighted "
          f"table runs, {unhl} unhighlighted body runs")
    print(f"  marked copy: convention note is the first paragraph: {note_ok}")
    assert hl > 0 and unhl > 0 and note_ok
    mt = docx_paragraphs(dm)
    assert mt[0] == norm(MARK_NOTE), mt[0][:80]
    assert mt[1:] == got, "marked copy text differs from the clean copy"
    print("  OK  marked copy carries exactly the clean text plus the note")
    mb = bold_audit(dm)
    print(f"  run-level bold runs in the marked copy: {len(mb)}")

    probe = "R.T. designed the project"
    for p in dm.paragraphs:
        if probe in p.text:
            n_hl = sum(1 for r in p.runs
                       if r.font.highlight_color == WD_COLOR_INDEX.YELLOW)
            print(f"  unchanged v4 paragraph ('Author Contributions') "
                  f"highlighted runs: {n_hl}")
            assert n_hl == 0, "unchanged v4 text was highlighted"
            break

    # ---- PDFs ------------------------------------------------------------
    pages_clean, size_clean = pdf_check(OUT_CLEAN, "clean ")
    pages_marked, size_marked = pdf_check(OUT_MARKED, "marked")

    print("=" * 76)
    print(f"ALL CHECKS PASSED   clean {pages_clean} pages ({size_clean}), "
          f"marked {pages_marked} pages ({size_marked})")
    return pages_clean, pages_marked


if __name__ == "__main__":
    only = sys.argv[1] if len(sys.argv) > 1 else ""
    do_clean = only != "--marked-only"
    do_marked = only != "--clean-only"
    rc = rm = None
    try:
        if do_clean:
            rc = build(OUT_CLEAN, marked=False)
            print(f"wrote {OUT_CLEAN}")
        if do_marked:
            rm = build(OUT_MARKED, marked=True)
            print(f"wrote {OUT_MARKED}")
        if do_clean and do_marked:
            verify(rc, rm)
    finally:
        if _TRIM_DIR:
            shutil.rmtree(_TRIM_DIR, ignore_errors=True)
